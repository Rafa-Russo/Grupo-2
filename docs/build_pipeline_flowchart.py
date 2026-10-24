"""Gera docs/pipeline_sp.html: caderno passo a passo do pipeline SP, só com linhas reais.

Uso: .venv/Scripts/python docs/build_pipeline_flowchart.py

Nada é digitado à mão: cada número e cada linha mostrada vem dos parquets em data/SP/cache,
das planilhas brutas (varredura cacheada em data/SP/cache/_doc_bruto.json) ou é recalculado aqui
com a mesma lógica dos notebooks 01_preparacao e 02_eda. Asserts no meio do caminho garantem
que o recálculo bate com os parquets gravados pelos notebooks.
"""
import html
import json
import re
import unicodedata
from datetime import datetime
from pathlib import Path

import geopandas as gpd
import numpy as np
import pandas as pd
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
DADOS = RAIZ / "data" / "SP"
CACHE = DADOS / "cache"
SAIDA = RAIZ / "docs" / "pipeline_sp.html"
BRUTO_JSON = CACHE / "_doc_bruto.json"

RUA = "BARAO DE LADARIO"
SQL_A = "02500700607"
ANOS_COMPLETOS = [2022, 2023, 2024, 2025]
JANELA = {"lat": (-24.0, -23.3), "lon": (-46.9, -46.3)}

# ═════════════════════════ lógica espelhada dos notebooks ═════════════════════════
TIPOS_VIA = {
    "R", "RUA", "AV", "AVENIDA", "TV", "TRAVESSA", "PC", "PRACA", "AL", "ALAMEDA",
    "EST", "ESTR", "ESTRADA", "LG", "LARGO", "VD", "VIADUTO", "VIA", "ROD", "RODOVIA",
    "MARGINAL", "BC", "BECO", "PSG", "PASSAGEM", "PASSEIO", "PTE", "PONTE", "ACS",
    "ACESSO", "COR", "CORREDOR", "ESC", "ESCADARIA", "TR", "TREVO", "BALAO", "COMP",
    "COMPLEXO", "PRQ", "PARQUE", "PQ", "JD", "JARDIM", "CJ", "CONJUNTO", "SIT", "MKD",
}
ABREVIACOES = {
    "S": "SAO", "STO": "SANTO", "STA": "SANTA", "D": "DOM", "PE": "PADRE",
    "PROF": "PROFESSOR", "PROFA": "PROFESSORA", "DR": "DOUTOR", "DRA": "DOUTORA",
    "CEL": "CORONEL", "GEN": "GENERAL", "CAP": "CAPITAO", "ENG": "ENGENHEIRO",
    "SEN": "SENADOR", "DEP": "DEPUTADO", "PRES": "PRESIDENTE", "CDE": "CONDE",
    "VISC": "VISCONDE", "BRIG": "BRIGADEIRO", "MAL": "MARECHAL", "REV": "REVERENDO",
    "MIN": "MINISTRO", "DES": "DESEMBARGADOR", "SGT": "SARGENTO", "TEN": "TENENTE",
    "MAJ": "MAJOR", "ALM": "ALMIRANTE", "N": "NOSSA", "SRA": "SENHORA",
    "MONS": "MONSENHOR", "VER": "VEREADOR", "COM": "COMENDADOR", "BR": "BARAO",
    "MQS": "MARQUES", "IRM": "IRMA", "CONS": "CONSELHEIRO", "PRINC": "PRINCESA",
}
VAZIOS = {"", "NULL", "NAN", "NONE"}
ABREVIACOES_BAIRRO = {
    "VL": "VILA", "JD": "JARDIM", "PQ": "PARQUE", "CJ": "CONJUNTO", "CID": "CIDADE",
    "S": "SAO", "STO": "SANTO", "STA": "SANTA", "CH": "CHACARA", "AL": "ALTO",
}
MINIMO_OCORRENCIAS_RUA = 20
VIOLENTOS = ("ROUBO", "LATROCINIO", "HOMICIDIO", "LESAO CORPORAL", "ESTUPRO", "EXTORSAO", "SEQUESTRO")
CENTRO_COMERCIAL = ["SE", "REPUBLICA", "BOM RETIRO", "BRAS", "SANTA IFIGENIA", "LUZ", "CAMPOS ELISEOS"]

COLUNAS_ITBI = [
    "sql", "logradouro", "numero", "complemento", "bairro_itbi", "referencia", "cep",
    "natureza_transacao", "valor_transacao", "data_transacao", "vvr",
    "proporcao_transmitida", "vvr_proporcional", "base_calculo", "tipo_financiamento",
    "valor_financiado", "cartorio", "matricula", "situacao_sql", "area_terreno",
    "testada", "fracao_ideal", "area_construida", "uso_cod", "uso_desc", "padrao_cod",
    "padrao_desc", "acc",
]
NUMERICAS_ITBI = [
    "valor_transacao", "vvr", "vvr_proporcional", "base_calculo", "valor_financiado",
    "proporcao_transmitida", "area_terreno", "testada", "fracao_ideal",
    "area_construida", "acc", "uso_cod", "padrao_cod",
]
DESCARTAVEIS_ITBI = ["complemento", "referencia", "cartorio", "tipo_financiamento", "matricula"]
CAMPOS_CRIME = {
    "municipio": ("NOME_MUNICIPIO_CIRCUN",), "data_ocorrencia": ("DATA_OCORRENCIA_BO",),
    "periodo": ("DESC_PERIODO", "DESCR_PERIODO"), "local": ("DESCR_SUBTIPOLOCAL",),
    "bairro": ("BAIRRO",), "logradouro": ("LOGRADOURO",), "lat": ("LATITUDE",),
    "lon": ("LONGITUDE",), "natureza": ("NATUREZA_APURADA",), "delegacia": ("NOME_DELEGACIA_CIRCUN",),
}


def sem_acento(texto):
    return unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode().upper().strip()


def limpa_texto(valor):
    texto = sem_acento(valor)
    return None if texto in VAZIOS else texto


def chave_rua(nome):
    texto = limpa_texto(nome)
    if texto is None or "VEDACAO" in texto:
        return None
    palavras = re.sub(r"[^A-Z0-9 ]", " ", texto).split()
    while palavras and palavras[0] in TIPOS_VIA:
        palavras = palavras[1:]
    return " ".join(ABREVIACOES.get(p, p) for p in palavras) or None


def em_unicos(serie, funcao):
    return serie.map({v: funcao(v) for v in serie.dropna().unique()})


def normaliza_bairro(nome):
    texto = unicodedata.normalize("NFKD", str(nome)).encode("ascii", "ignore").decode().upper()
    palavras = re.sub(r"[^A-Z0-9 ]", " ", texto).split()
    return " ".join(ABREVIACOES_BAIRRO.get(p, p) for p in palavras) or None


def grupo_crime(natureza):
    if "ACIDENTE DE TRANSITO" in natureza or "CULPOS" in natureza:
        return "trânsito / culposo"
    if natureza.startswith("FURTO"):
        return "patrimonial sem violência"
    if any(t in natureza for t in VIOLENTOS):
        return "violento"
    return "outros"


def limpa_crimes_linhas(t):
    """Mesma limpeza de limpa_crimes(), sem o filtro final de ano (o filtro é mostrado à parte)."""
    t = t.copy()
    for c in ["bairro", "natureza", "periodo", "local", "delegacia"]:
        if c in t:
            t[c] = em_unicos(t[c], limpa_texto)
    t["logradouro"] = em_unicos(t["logradouro"], limpa_texto)
    t["rua_key"] = em_unicos(t["logradouro"], chave_rua)
    t["data_ocorrencia"] = pd.to_datetime(t["data_ocorrencia"], errors="coerce")
    t["ano"] = t["data_ocorrencia"].dt.year
    t["mes"] = t["data_ocorrencia"].dt.to_period("M").astype(str)
    for c, (lo, hi) in JANELA.items():
        v = pd.to_numeric(t[c], errors="coerce")
        t[c] = v.where(v.between(lo, hi))
    return t


def limpa_itbi_linhas(t):
    t = t.copy()
    t[NUMERICAS_ITBI] = t[NUMERICAS_ITBI].apply(pd.to_numeric, errors="coerce")
    t["sql"] = t["sql"].astype("int64").astype(str).str.zfill(11)
    t["setor_fiscal"] = t["sql"].str[:3]
    t["logradouro"] = em_unicos(t["logradouro"], limpa_texto)
    t["rua_key"] = em_unicos(t["logradouro"], chave_rua)
    t["bairro_itbi"] = em_unicos(t["bairro_itbi"], limpa_texto)
    t["data_transacao"] = pd.to_datetime(t["data_transacao"], errors="coerce")
    t["ano"] = t["data_transacao"].dt.year
    t["mes"] = t["data_transacao"].dt.to_period("M").astype(str)
    for c in ["natureza_transacao", "uso_desc", "padrao_desc", "setor_fiscal", "situacao_sql"]:
        t[c] = em_unicos(t[c], limpa_texto)
    return t.drop(columns=DESCARTAVEIS_ITBI)


# ═════════════════════════ formatação HTML ═════════════════════════
def e(x):
    return html.escape(str(x))


def fm(n):
    return f"{int(round(n)):,}".replace(",", ".")


def pc(x, d=1):
    return f"{x * 100:.{d}f}".replace(".", ",") + "%"


def num(x, d=2):
    return f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def brl(x, d=0):
    return "R$ " + num(x, d)


class H(str):
    """Texto que já é HTML (não escapar)."""


def nulo(v):
    return v is None or (not isinstance(v, (str, bytes)) and bool(pd.isna(v)))


def fmt(v):
    """Valor de célula já tipado (pós-limpeza)."""
    if nulo(v):
        return '<span class="nulo">∅</span>'
    if isinstance(v, (pd.Timestamp, datetime)):
        return v.strftime("%Y-%m-%d") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, H) or (isinstance(v, str) and '<span class="' in v):
        return str(v)
    if isinstance(v, (bool, np.bool_)):
        return "sim" if v else "não"
    if isinstance(v, (int, np.integer)):
        return str(int(v))
    if isinstance(v, (float, np.floating)):
        if float(v).is_integer() and abs(v) < 1e15:
            return str(int(v))
        t = f"{v:.7f}" if abs(v) < 1000 else f"{v:,.2f}"
        t = t.rstrip("0").rstrip(".")
        return t.replace(",", "X").replace(".", ",").replace("X", ".")
    return e(v)


def bruto(v):
    """Valor de célula como veio da planilha: mostra o tipo (aspas = texto, .0 = float)."""
    return H(_bruto(v))


def _bruto(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return '<span class="nulo">∅ vazio</span>'
    if isinstance(v, str):
        return f'<span class="txt">"{e(v)}"</span>'
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, np.integer)):
        return str(int(v))
    if isinstance(v, (float, np.floating)):
        return repr(float(v))
    return e(v)


def norm(v):
    if nulo(v):
        return None
    if isinstance(v, (pd.Timestamp, datetime)):
        return pd.Timestamp(v).strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, (int, float, np.integer, np.floating)):
        return f"{float(v):.6g}"
    return str(v)


def render(df, marks=None, colcls=None, rowcls=None, labels=None, transposed=None, texto_bruto=False):
    """Tabela. marks[(i,col)], colcls[col], rowcls[i] → classes css. Transposta se tiver muitas colunas."""
    marks, colcls, rowcls = marks or {}, colcls or {}, rowcls or {}
    df = df.reset_index(drop=True)
    cell = bruto if texto_bruto else fmt
    if transposed is None:
        transposed = df.shape[1] > 7
    labels = labels or [f"linha {i + 1}" for i in range(len(df))]
    out = ['<div class="rolagem"><table class="campos">' if transposed else '<div class="rolagem"><table class="largo">']
    if transposed:
        out.append("<thead><tr><th>campo</th>" + "".join(f"<th>{e(lab)}</th>" for lab in labels) + "</tr></thead><tbody>")
        for c in df.columns:
            cc = colcls.get(c, "")
            out.append(f'<tr class="{cc}"><th>{e(c)}</th>')
            for i, v in enumerate(df[c]):
                k = marks.get((i, c)) or rowcls.get(i) or ""
                out.append(f'<td class="{k}">{cell(v)}</td>')
            out.append("</tr>")
    else:
        out.append("<thead><tr>" + "".join(f'<th class="{colcls.get(c, "")}">{e(c)}</th>' for c in df.columns) + "</tr></thead><tbody>")
        for i in range(len(df)):
            out.append(f'<tr class="{rowcls.get(i, "")}">')
            for c in df.columns:
                k = marks.get((i, c)) or colcls.get(c) or ""
                out.append(f'<td class="{k}">{cell(df.iloc[i][c])}</td>')
            out.append("</tr>")
    out.append("</tbody></table></div>")
    return "".join(out)


def diff_tabela(antes, depois, labels=None, antes_bruto=True, ocultar=()):
    """Campo × registros; cada registro ocupa duas colunas: antes | depois. Colorida pelo que mudou."""
    antes, depois = antes.reset_index(drop=True), depois.reset_index(drop=True)
    labels = labels or [f"linha {i + 1}" for i in range(len(antes))]
    campos = [c for c in antes.columns if c not in ocultar] + [c for c in depois.columns if c not in antes.columns and c not in ocultar]
    cab = "".join(f'<th colspan="2" class="reg">{e(lab)}</th>' for lab in labels)
    sub = "".join("<th>antes</th><th>depois</th>" for _ in labels)
    out = [f'<div class="rolagem"><table class="campos diff"><thead><tr><th rowspan="2">campo</th>{cab}</tr><tr>{sub}</tr></thead><tbody>']
    fa = bruto if antes_bruto else fmt
    for c in campos:
        out.append(f"<tr><th>{e(c)}</th>")
        for i in range(len(antes)):
            em_a, em_d = c in antes.columns, c in depois.columns
            va = antes.at[i, c] if em_a else None
            vd = depois.at[i, c] if em_d else None
            if em_a and not em_d:
                out.append(f'<td class="cai">{fa(va)}</td><td class="vazio">—</td>')
            elif em_d and not em_a:
                out.append(f'<td class="vazio">—</td><td class="novo">{fmt(vd)}</td>')
            else:
                mudou = norm(va) != norm(vd)
                out.append(f"<td>{fa(va)}</td><td class=\"{'mod' if mudou else ''}\">{fmt(vd)}</td>")
        out.append("</tr>")
    out.append("</tbody></table></div>")
    return "".join(out)


def codigo(txt):
    return f'<pre class="cod">{e(txt)}</pre>'


def caso(titulo, regra, n_txt, corpo, tipo="trata"):
    selo = {"trata": "tratado", "limite": "NÃO tratado", "info": "contexto"}[tipo]
    return (
        f'<div class="caso {tipo}"><div class="caso-cab"><span class="selo">{selo}</span><h4>{titulo}</h4>'
        f'<span class="n">{n_txt}</span></div><p class="regra">{regra}</p>{corpo}</div>'
    )


def passo(pid, numero, titulo, intro, blocos):
    return (
        f'<section class="passo" id="{pid}"><h2><span>{numero}</span>{titulo}</h2>'
        f'<p class="intro">{intro}</p>{"".join(blocos)}</section>'
    )


def sub(titulo, corpo, nota=""):
    return f'<div class="sub"><h3>{titulo}</h3>{f"<p class=nota>{nota}</p>" if nota else ""}{corpo}</div>'


def barra_tabela(linhas, cab=("etapa", "linhas", "", "%")):
    """linhas: [(nome, n, total_ref)] → tabela com barra."""
    corpo = "".join(
        f'<tr><td>{n}</td><td class="num">{fm(q)}</td><td class="barra"><i style="width:{min(q / ref * 100, 100):.2f}%"></i></td>'
        f'<td class="num">{pc(q / ref, 2)}</td></tr>'
        for n, q, ref in linhas
    )
    ch = "".join(f"<th>{c}</th>" for c in cab)
    return f'<div class="rolagem"><table class="largo"><thead><tr>{ch}</tr></thead><tbody>{corpo}</tbody></table></div>'


# ═════════════════════════ carga dos caches ═════════════════════════
print("lendo caches…")
crimes = pd.read_parquet(CACHE / "crimes_sp.parquet")
itbi = pd.read_parquet(CACHE / "itbi_sp.parquet")
dic = pd.read_parquet(CACHE / "dicionario_ruas.parquet")
setor_dist = pd.read_parquet(CACHE / "setor_para_distrito.parquet")
ipca_csv = pd.read_csv(CACHE / "ipca_ibge.csv", dtype={"periodo": str})
bairros = pd.read_parquet(CACHE / "bairros.parquet")
base_pq = pd.read_parquet(CACHE / "base_modelagem.parquet")
MES_REF = (CACHE / "mes_referencia.txt").read_text(encoding="utf-8").strip()
distritos = gpd.read_file(CACHE / "distritos_sp.geojson")
setores_raw = gpd.read_file(CACHE / "setores_fiscais.geojson")

# ───────── IPCA (mesma lógica de carrega_ipca) ─────────
ipca = ipca_csv.copy()
for c in ("indice", "variacao_pct"):
    ipca[c] = pd.to_numeric(ipca[c], errors="coerce")
ipca["mes"] = ipca["periodo"].str[:4] + "-" + ipca["periodo"].str[4:]
ipca = ipca.dropna(subset=["indice"]).sort_values("mes").reset_index(drop=True)
ipca["fator"] = ipca["indice"].iloc[-1] / ipca["indice"]
assert ipca["mes"].iloc[-1] == MES_REF

# ───────── crime → distrito (sjoin da seção 3 do notebook 01) ─────────
print("sjoin crimes × distritos…")
geoloc = crimes.dropna(subset=["lat", "lon"])
pontos = gpd.GeoDataFrame({"indice": geoloc.index}, geometry=gpd.points_from_xy(geoloc["lon"], geoloc["lat"]), crs=4326)
dentro = gpd.sjoin(pontos, distritos[["nm_distrito_municipal", "geometry"]], how="left", predicate="within")
n_pontos, n_linhas_sjoin = len(pontos), len(dentro)
dup_idx = dentro["indice"][dentro["indice"].duplicated(keep=False)]
n_fronteira = dup_idx.nunique()
n_fora = int(dentro.drop_duplicates("indice")["nm_distrito_municipal"].isna().sum())
exemplo_fronteira = dentro[dentro["indice"] == dup_idx.iloc[0]] if len(dup_idx) else None
mapa_distrito = dentro.drop_duplicates("indice").set_index("indice")["nm_distrito_municipal"]
crimes["distrito"] = crimes.index.map(mapa_distrito)

# ───────── localiza() espelhada, mantendo todas as colunas intermediárias ─────────
print("localiza()…")
VOCAB = set(crimes["bairro"].dropna().astype(str).unique())


def localiza_completo(transacoes):
    t = transacoes.copy()
    t["setor_fiscal"] = t["setor_fiscal"].astype(str)
    traduzido = {v: normaliza_bairro(v) for v in t["bairro_itbi"].dropna().unique()}
    t["cep_txt"] = t["cep"].astype("Int64").astype(str).str.zfill(8)
    t["bairro_norm"] = t["bairro_itbi"].map(traduzido)
    t["bairro_declarado"] = t["bairro_norm"].where(t["bairro_norm"].isin(VOCAB))
    n0 = len(t)
    t = t.merge(setor_dist[["setor_fiscal", "distrito"]], on="setor_fiscal", how="left")
    assert len(t) == n0
    votos = (
        t.dropna(subset=["bairro_declarado"]).groupby(["cep_txt", "bairro_declarado"]).size().rename("n").reset_index()
        .sort_values("n").groupby("cep_txt")["bairro_declarado"].last().rename("bairro_cep")
    )
    t = t.join(votos, on="cep_txt")
    conf = dic[dic["n_ocorrencias"].ge(MINIMO_OCORRENCIAS_RUA)]
    por_dist = conf[conf["distrito"].notna()].set_index(["rua_key", "distrito"])
    cidade = conf[conf["distrito"].isna()].set_index("rua_key")
    t = t.join(por_dist[["bairro", "lat", "lon"]].add_suffix("_local"), on=["rua_key", "distrito"])
    t = t.join(cidade[["bairro", "lat", "lon"]].add_suffix("_cidade"), on="rua_key")
    t["bairro_rua"] = t["bairro_local"].fillna(t["bairro_cidade"])
    t["bairro"] = t["bairro_declarado"].fillna(t["bairro_cep"]).fillna(t["bairro_rua"])
    t["fonte"] = np.select(
        [t["bairro_declarado"].notna(), t["bairro_cep"].notna(), t["bairro_local"].notna(), t["bairro_cidade"].notna()],
        ["1 · bairro declarado", "2 · CEP", "3a · rua + distrito", "3b · rua cidade-toda"], "sem bairro",
    )
    centro = crimes.dropna(subset=["lat", "lon"]).groupby("bairro", observed=True)[["lat", "lon"]].median()
    centro.index = centro.index.astype(str)
    t["rua_concorda"] = t["bairro_rua"].eq(t["bairro"])
    for eixo in ["lat", "lon"]:
        t[f"{eixo}_rua"] = t[f"{eixo}_local"].fillna(t[f"{eixo}_cidade"])
        t[f"{eixo}_centro"] = t["bairro"].map(centro[eixo])
        t[eixo] = t[f"{eixo}_rua"].where(t["rua_concorda"], t[f"{eixo}_centro"])
    return t


loc = localiza_completo(itbi)
assert len(loc) == len(itbi)

# ───────── filtros e preço (seção 4 do notebook 02) ─────────
MESES_IPCA = set(ipca["mes"])
FILTROS = [
    ("transação entre 2022 e 2026", lambda d: d["ano"].between(2022, 2026), ["ano", "data_transacao"]),
    ("natureza = compra e venda", lambda d: d["natureza_transacao"].astype(str).str.startswith("1."), ["natureza_transacao"]),
    ("100% do imóvel transmitido", lambda d: d["proporcao_transmitida"] == 100, ["proporcao_transmitida", "fracao_ideal"]),
    ("padrão residencial", lambda d: d["padrao_desc"].astype(str).str.startswith("RESIDENCIAL"), ["padrao_desc"]),
    ("exclui vaga de garagem", lambda d: ~d["uso_desc"].astype(str).str.contains("GARAGEM"), ["uso_desc"]),
    ("área construída e valor > 0", lambda d: d["area_construida"].gt(0) & d["valor_transacao"].gt(0), ["area_construida", "valor_transacao", "uso_desc"]),
    ("bairro identificado", lambda d: d["bairro"].notna(), ["bairro_itbi", "cep", "rua_key"]),
    ("mês coberto pelo IPCA", lambda d: d["mes"].isin(MESES_IPCA), ["mes", "data_transacao"]),
]
funil, removidas, atual = [("guias de ITBI (tudo)", len(loc))], {}, loc
for nome, f, _ in FILTROS:
    m = f(atual)
    removidas[nome] = atual[~m]
    atual = atual[m]
    funil.append((nome, len(atual)))
base = atual.merge(ipca[["mes", "fator"]], on="mes", how="left")
base["valor_real"] = base["valor_transacao"] * base["fator"]
base["preco_m2_nominal"] = base["valor_transacao"] / base["area_construida"]
base["preco_m2"] = base["valor_real"] / base["area_construida"]
P1, P99 = base["preco_m2"].quantile([0.01, 0.99])
n_antes_pct = len(base)
fora_pct = base[~base["preco_m2"].between(P1, P99)]
base = base[base["preco_m2"].between(P1, P99)].copy()
funil.append(("R$/m² real entre os percentis 1 e 99", len(base)))
base["idade"] = base["ano"] - base["acc"].where(base["acc"].between(1800, 2026))
base["tipo"] = base["padrao_desc"].astype(str).str.strip()

# ───────── crime → grupo → bairro (seções 2 e 8) ─────────
crimes["grupo"] = crimes["natureza"].astype(str).map(grupo_crime)
por_bairro_grupo_ano = (
    crimes[crimes["ano"].isin(ANOS_COMPLETOS)].groupby(["bairro", "grupo"], observed=True).size().unstack(fill_value=0)
)
por_bairro_grupo = por_bairro_grupo_ano.div(len(ANOS_COMPLETOS))
por_bairro_grupo.index = por_bairro_grupo.index.astype(str)
preco_bairro = base.groupby("bairro").agg(
    preco_m2_mediano=("preco_m2", "median"), n_transacoes=("preco_m2", "size"),
    area_mediana=("area_construida", "median"), idade_mediana=("idade", "median"),
)
bair = preco_bairro.join(por_bairro_grupo, how="inner")
bair["crimes_total"] = bair[por_bairro_grupo.columns].sum(axis=1)
bair["violentos"] = bair["violento"]
bair["share_violento"] = bair["violentos"] / bair["crimes_total"]
bair["violentos_por_1k"] = bair["violentos"] / bair["n_transacoes"] * 1000
bair_aprovados = bair[bair["n_transacoes"].ge(30) & bair["crimes_total"].ge(30)]
COLS_CRIME = ["crimes_total", "violentos", "share_violento", "violentos_por_1k"]
modelagem = base.merge(bair_aprovados[COLS_CRIME], left_on="bairro", right_index=True, how="inner")

# o recálculo precisa reproduzir exatamente o que os notebooks gravaram
assert len(modelagem) == len(base_pq), (len(modelagem), len(base_pq))
assert np.allclose(np.sort(modelagem["preco_m2"].to_numpy()), np.sort(base_pq["preco_m2"].to_numpy()))
assert len(bair_aprovados) == len(bairros)
print(f"recálculo confere: base {len(modelagem):,} linhas, {len(bair_aprovados)} bairros")

crimes_com_rua_bairro = crimes.dropna(subset=["rua_key", "bairro"])


# ═════════════════════════ varredura das planilhas brutas (cacheada) ═════════════════════════
def escolhe_alvos():
    """SQLs cujas linhas brutas queremos mostrar."""
    alvos = {}
    ok = modelagem
    alvos["A"] = SQL_A
    pool_b = ok[(ok["bairro_itbi"] == "STO AMARO") & (ok["fonte"] == "1 · bairro declarado")].sort_values("sql")
    alvos["B"] = pool_b["sql"].iloc[0]
    pool_c = ok[ok["fonte"] == "3a · rua + distrito"].sort_values("sql")
    alvos["C"] = pool_c["sql"].iloc[len(pool_c) // 2]
    r3 = removidas["100% do imóvel transmitido"]
    r3 = r3[(r3["bairro"].notna()) & (r3["proporcao_transmitida"] < 100) & (r3["ano"] >= 2024)].sort_values("sql")
    alvos["D"] = r3["sql"].iloc[len(r3) // 3]
    for ano, k in [(2022, "L22"), (2024, "L24"), (2026, "L26")]:
        p = ok[(ok["ano"] == ano) & (ok["fonte"] == "2 · CEP")].sort_values("sql")
        alvos[k] = p["sql"].iloc[len(p) // 2]
    return alvos


ALVOS = escolhe_alvos()
ALVOS_INT = {int(v) for v in ALVOS.values()}


def varre_brutos():
    print("varrendo planilhas brutas (≈ 2 min na primeira vez)…")
    res = {"versao": 2, "alvos": sorted(ALVOS_INT), "crime_abas": [], "crime_ex": {}, "crime_cont": {}, "itbi_arquivos": [], "itbi_alvos": {}}
    ex, cont = res["crime_ex"], res["crime_cont"]
    n_ex = 3

    def guarda(classe, linha, aceita=True):
        cont[classe] = cont.get(classe, 0) + 1
        lst = ex.setdefault(classe, [])
        chave = (linha["logradouro"], linha["bairro"])
        if aceita and len(lst) < n_ex and all((x["logradouro"], x["bairro"]) != chave for x in lst):
            lst.append(linha)

    for caminho in sorted((DADOS / "Crimes").glob("*.xlsx")):
        wb = load_workbook(caminho, read_only=True)
        for aba in [a for a in wb.sheetnames if re.fullmatch(r"(JAN-JUN|JUL-DEZ)_\d{4}", a)]:
            it = wb[aba].iter_rows(values_only=True)
            cab_cru = list(next(it))
            cab = [sem_acento(c) for c in cab_cru]
            pos = {k: next(i for i, c in enumerate(cab) if c.startswith(pref)) for k, pref in CAMPOS_CRIME.items()}
            casou = {k: cab_cru[i] for k, i in pos.items()}
            total = sp = 0
            for linha in it:
                total += 1
                if linha[pos["municipio"]] != "S.PAULO":
                    continue
                sp += 1
                reg = {k: linha[i] for k, i in pos.items()}
                reg = {k: (v.isoformat(sep=" ") if isinstance(v, datetime) else v) for k, v in reg.items()}
                reg["_data_tipo"] = type(linha[pos["data_ocorrencia"]]).__name__
                reg["_aba"] = aba
                d = linha[pos["data_ocorrencia"]]
                vedado = "VEDACAO" in sem_acento(linha[pos["logradouro"]] or "")
                ano_ok = isinstance(d, datetime) and 2022 <= d.year <= 2026
                limpo = ano_ok and not vedado
                lat = linha[pos["lat"]]
                if lat is None:
                    guarda("lat_vazia", reg, limpo)
                else:
                    try:
                        x = float(lat)
                        if x == 0:
                            guarda("lat_zero", reg, limpo)
                        elif abs(x) > 1000:
                            guarda("lat_sem_virgula", reg, limpo)
                        elif not JANELA["lat"][0] <= x <= JANELA["lat"][1]:
                            guarda("lat_fora", reg, limpo)
                        else:
                            guarda("lat_ok", reg, limpo)
                    except (TypeError, ValueError):
                        guarda("lat_texto", reg, limpo)
                if not isinstance(d, datetime):
                    guarda("data_invalida", reg, not vedado)
                elif not 2022 <= d.year <= 2026:
                    guarda("ano_fora", reg, not vedado)
                if vedado:
                    guarda("logradouro_vedado", reg, ano_ok)
                txt = f"{linha[pos['bairro']] or ''}{linha[pos['logradouro']] or ''}"
                if txt != sem_acento(txt) or any(ch.islower() for ch in txt):
                    guarda("acento", reg, limpo)
            res["crime_abas"].append({
                "arquivo": caminho.name, "aba": aba, "n_colunas": len(cab_cru), "casou": casou, "total": total, "sp": sp,
            })
            print(f"  {aba}: {sp:,}/{total:,}")
        wb.close()

    for caminho in sorted((DADOS / "Imoveis").glob("*.xlsx")):
        wb = load_workbook(caminho, read_only=True)
        meses = [a for a in wb.sheetnames if re.fullmatch(r"[A-Z]{3}-\d{4}", a)]
        info = {"arquivo": caminho.name, "meses": meses, "ignoradas": [a for a in wb.sheetnames if a not in meses],
                "mantidas": 0, "descartadas": 0, "descartes": [], "cab_topo": None, "primeira_celula": None}
        for k, aba in enumerate(meses):
            linhas = list(wb[aba].iter_rows(values_only=True))
            if k == 0:
                info["primeira_celula"] = linhas[0][0]
                info["cab_topo"] = not isinstance(linhas[0][0], (int, float))
            for i, linha in enumerate(linhas):
                if isinstance(linha[0], (int, float)):
                    info["mantidas"] += 1
                    if int(linha[0]) in ALVOS_INT:
                        res["itbi_alvos"].setdefault(str(int(linha[0])), []).append(
                            {"arquivo": caminho.name, "aba": aba, "linha": [v.isoformat(sep=" ") if isinstance(v, datetime) else v for v in (list(linha) + [None] * 28)[:28]]}
                        )
                else:
                    info["descartadas"] += 1
                    if len(info["descartes"]) < 3:
                        pos_txt = "topo" if i == 0 else ("fim" if i == len(linhas) - 1 else f"linha {i + 1}")
                        info["descartes"].append({"aba": aba, "onde": pos_txt, "celula": linha[0]})
        wb.close()
        res["itbi_arquivos"].append(info)
        print(f"  {caminho.name}: {info['mantidas']:,} mantidas")
    return res


RAW = json.loads(BRUTO_JSON.read_text(encoding="utf-8")) if BRUTO_JSON.exists() else {}
if RAW.get("alvos") != sorted(ALVOS_INT) or RAW.get("versao") != 2:
    RAW = varre_brutos()
    BRUTO_JSON.write_text(json.dumps(RAW, ensure_ascii=False, default=str), encoding="utf-8")


def linha_bruta_itbi(sql, valor_ref=None):
    cand = RAW["itbi_alvos"][str(int(sql))]
    if valor_ref is not None:
        cand = sorted(cand, key=lambda c: abs((c["linha"][8] or 0) - valor_ref))
    return cand[0]


def df_bruto_itbi(sqls):
    linhas, fontes = [], []
    for s in sqls:
        v = base_ou_loc(s)["valor_transacao"]
        c = linha_bruta_itbi(s, v)
        linhas.append(c["linha"])
        fontes.append(f'{c["arquivo"]} · {c["aba"]}')
    return pd.DataFrame(linhas, columns=COLUNAS_ITBI, dtype=object), fontes


def base_ou_loc(sql):
    return loc[loc["sql"] == sql].iloc[0]


# ═════════════════════════ SEÇÕES ═════════════════════════
secoes = []
TOTAL_ITBI, TOTAL_CRIMES = len(itbi), len(crimes)

# ───────────────────────── 0 · rastreio ─────────────────────────
def linha_do_tempo(passos):
    out = ['<div class="rolagem"><table class="tempo"><tbody>']
    for titulo, pares in passos:
        cel = "".join(
            f'<span class="par {cls}"><b>{e(k)}</b> {v}</span>' for k, v, cls in pares
        )
        out.append(f'<tr><th>{titulo}</th><td>{cel}</td></tr>')
    out.append("</tbody></table></div>")
    return "".join(out)


def par(k, v, cls="", bruto_=False):
    return (k, bruto(v) if bruto_ else fmt(v), cls)


def rastreia(chave, titulo, descricao):
    sql = ALVOS[chave]
    r = loc[loc["sql"] == sql].iloc[0]
    c = linha_bruta_itbi(sql, r["valor_transacao"])
    raw = dict(zip(COLUNAS_ITBI, c["linha"]))
    passos = [
        (f"1 · célula bruta<br><small>{e(c['arquivo'])} · {e(c['aba'])}</small>", [
            par("sql", raw["sql"], bruto_=True), par("logradouro", raw["logradouro"], bruto_=True),
            par("número", raw["numero"], bruto_=True), par("bairro", raw["bairro_itbi"], bruto_=True),
            par("cep", raw["cep"], bruto_=True), par("valor", raw["valor_transacao"], bruto_=True),
            par("data", raw["data_transacao"], bruto_=True), par("área", raw["area_construida"], bruto_=True),
        ]),
        ("2 · limpeza (notebook 01)", [
            par("sql", r["sql"], "mod" if str(raw["sql"]) != r["sql"] else ""),
            par("setor_fiscal", r["setor_fiscal"], "novo"), par("rua_key", r["rua_key"], "novo"),
            par("bairro_itbi", r["bairro_itbi"], "mod" if norm(raw["bairro_itbi"]) != norm(r["bairro_itbi"]) else ""),
            par("mes", r["mes"], "novo"),
        ]),
        ("3 · setor → distrito", [par("setor_fiscal", r["setor_fiscal"]), par("distrito", r["distrito"], "novo")]),
    ]
    dist_txt = []
    if pd.notna(r["bairro_declarado"]):
        dist_txt.append(par("bairro_declarado", r["bairro_declarado"], "novo"))
    else:
        dist_txt.append(par("bairro_declarado", None))
    passos.append(("4 · localiza() → bairro", dist_txt + [
        par("bairro_cep", r["bairro_cep"], "novo"),
        par("bairro_rua", r["bairro_rua"], "novo"),
        par("fonte que decidiu", r["fonte"], "mod"),
        par("bairro", r["bairro"], "novo"),
        par("lat, lon", f"{r['lat']:.5f}, {r['lon']:.5f}" if pd.notna(r["lat"]) else None, "novo"),
    ]))
    # filtros
    removido_em = next((n for n, d in removidas.items() if sql in set(d["sql"])), None)
    if removido_em is not None:
        passos.append(("5 · filtros", [par("resultado", f"REMOVIDA em «{removido_em}»", "cai")] + [
            par(col, r[col]) for col in dict((n, cols) for n, _, cols in FILTROS)[removido_em]
        ]))
    else:
        b = modelagem[modelagem["sql"] == sql].iloc[0]
        passos.append(("5 · filtros", [par("resultado", "passou nos 8 filtros", "")]))
        passos.append(("6 · IPCA", [
            par("mes", b["mes"]), par("fator", b["fator"], "novo"),
            par("valor_real", b["valor_real"], "novo"), par("preco_m2", b["preco_m2"], "novo"),
        ]))
        passos.append(("7 · perfil do bairro", [
            par("bairro", b["bairro"]), par("crimes_total", b["crimes_total"], "novo"),
            par("violentos", b["violentos"], "novo"), par("share_violento", b["share_violento"], "novo"),
            par("violentos_por_1k", b["violentos_por_1k"], "novo"),
        ]))
    return caso(titulo, descricao, f"sql {sql}", linha_do_tempo(passos), "info")


def decisao(sql):
    return loc[loc["sql"] == sql].iloc[0]["fonte"]


assert decisao(SQL_A) == "2 · CEP" and SQL_A in set(modelagem["sql"])
secoes.append(passo(
    "s0", "0", "Rastreio: quatro transações do arquivo bruto até a base final",
    "Antes de entrar nos passos, quatro guias reais seguidas ponta a ponta — uma para cada destino possível. "
    "Os passos 1–10 abaixo explicam cada seta desta linha do tempo. "
    '<span class="par mod"><b>amarelo</b></span> = valor alterado · <span class="par novo"><b>verde</b></span> = campo criado · '
    '<span class="par cai"><b>vermelho</b></span> = descartado.',
    [
        rastreia("A", "A · bairro decidido pelo CEP", "Guia sem bairro válido; os vizinhos do mesmo CEP votam."),
        rastreia("B", "B · bairro declarado, abreviação expandida", "O campo Bairro vem abreviado e passa a valer depois de normalizado."),
        rastreia("C", "C · bairro decidido por rua + distrito", "Sem bairro e sem voto de CEP: o dicionário de ruas, na camada do distrito do setor fiscal, responde."),
        rastreia("D", "D · descartada no filtro de proporção", "Fração de um lote-matriz: a área construída seria do prédio inteiro, não do imóvel vendido."),
    ],
))

# ───────────────────────── 1 · crimes ─────────────────────────
abas = pd.DataFrame(RAW["crime_abas"])
abas["muni"] = abas["casou"].map(lambda d: d["municipio"])
abas["per"] = abas["casou"].map(lambda d: d["periodo"])
tab_abas = pd.DataFrame({
    "aba": abas["aba"], "colunas": abas["n_colunas"], "município ← coluna": abas["muni"],
    "período ← coluna": abas["per"], "linhas lidas": abas["total"].map(fm), "S.PAULO": abas["sp"].map(fm),
    "mantidas": (abas["sp"] / abas["total"]).map(pc),
})
assert abas["sp"].sum() >= TOTAL_CRIMES
n_anos_fora = RAW["crime_cont"].get("ano_fora", 0)
bloco_cab = caso(
    "Cabeçalho muda de ano para ano",
    "As colunas são achadas por <b>prefixo</b> (<code>NOME_MUNICIPIO_CIRCUN</code>, <code>DESC_PERIODO</code>/<code>DESCR_PERIODO</code>…), não por nome exato. "
    "Sem acento antes de comparar. Se nenhum prefixo casa, <code>acha_coluna</code> levanta erro em vez de seguir com coluna errada.",
    f"{len(abas)} abas lidas",
    render(tab_abas, transposed=False)
    + f'<p class="nota">Linhas lidas = estado inteiro; só <code>S.PAULO</code> é mantida já na leitura. Total mantido: <b>{fm(abas["sp"].sum())}</b> de {fm(abas["total"].sum())} '
    f'({pc(abas["sp"].sum() / abas["total"].sum())}). Depois do filtro de ano (2022–2026): <b>{fm(TOTAL_CRIMES)}</b> linhas em <code>crimes_sp.parquet</code>.</p>',
)


def exemplo_crime(classe, n=1):
    linhas = RAW["crime_ex"].get(classe, [])[:n]
    if not linhas:
        return None
    df = pd.DataFrame(linhas)
    cols = ["data_ocorrencia", "periodo", "local", "bairro", "logradouro", "lat", "lon", "natureza"]
    return df[cols], df["_aba"].tolist()


def bloco_crime(classe, titulo, regra, n=2, ocultar=("delegacia",), tipo="trata"):
    r = exemplo_crime(classe, n)
    if r is None:
        return ""
    antes, abas_ = r
    depois = limpa_crimes_linhas(antes)
    marc_rows = {}
    corpo = diff_tabela(antes, depois, labels=[f"{a}" for a in abas_], ocultar=ocultar)
    if classe == "ano_fora":
        corpo += '<p class="nota">Estas linhas são descartadas por <code>tabela["ano"].between(2022, 2026)</code> e não chegam ao parquet.</p>'
    return caso(titulo, regra, f"{fm(RAW['crime_cont'].get(classe, 0))} linhas", corpo, tipo)


blocos_crime = [bloco_cab]
blocos_crime.append(bloco_crime(
    "lat_texto", "LATITUDE = 'NULL' (texto no meio de números)",
    "<code>pd.to_numeric(errors='coerce')</code> transforma o texto em ∅; depois só valores dentro da janela da cidade sobrevivem.",
))
blocos_crime.append(bloco_crime(
    "lat_zero", "LATITUDE = 0",
    "Zero cai fora da janela (lat −24,0 a −23,3 · lon −46,9 a −46,3) e vira ∅. A mesma janela vale para longitude.",
))
blocos_crime.append(bloco_crime(
    "lat_sem_virgula", "Coordenada sem separador decimal",
    "A planilha traz <code>-2354897372</code> em vez de <code>-23.54897372</code>: o número inteiro, ×10⁸. Cai fora da janela e vira ∅. Dado recuperável que o pipeline descarta (ver passo 11).", tipo="limite",
))
blocos_crime.append(bloco_crime(
    "lat_fora", "Coordenada com formato válido, mas fora da janela de SP",
    "Pontos em outro município (ou com lat/lon trocados) não podem ser atribuídos a um distrito da capital: viram ∅ nos dois eixos.",
))
blocos_crime.append(bloco_crime(
    "lat_ok", "Coordenada boa (para comparar)",
    "O valor passa intacto; só o formato de texto muda. O ponto será usado no <code>sjoin</code> com os distritos (passo 3).", n=1,
))
blocos_crime.append(bloco_crime(
    "acento", "Acentos e minúsculas",
    "<code>sem_acento</code>: NFKD → ASCII → MAIÚSCULAS → trim. É isso que faz <code>BRÁS</code> e <code>BRAS</code> serem o mesmo bairro.",
))
blocos_crime.append(bloco_crime(
    "logradouro_vedado", "Logradouro com divulgação vedada",
    "Quando o texto contém <code>VEDACAO</code> (violência sexual e outros casos sigilosos) o <code>rua_key</code> fica ∅: o crime existe e conta para o bairro, mas não alimenta o dicionário de ruas.",
))
blocos_crime.append(bloco_crime(
    "data_invalida", "Data ilegível",
    "Texto não-data (<code>\"NULL\"</code>) vira NaT e ano/mês ficam ∅; data com ano absurdo (<code>1202-10-15</code>) é lida, mas cai fora do filtro 2022–2026. Nos dois casos a linha é descartada.", tipo="trata",
))
blocos_crime.append(bloco_crime(
    "ano_fora", "Fato ocorrido antes de 2022 (registrado depois)",
    "O BO é de 2022+, mas a data do fato é anterior. O pipeline usa a data do <b>fato</b>.",
))

# chave_rua: exemplos reais, escolhidos por frequência
ruas_freq = crimes["logradouro"].value_counts()
lista = pd.DataFrame({"logradouro": ruas_freq.index, "n": ruas_freq.values})
lista["rua_key"] = lista["logradouro"].map(chave_rua)


def tipos_removidos(txt):
    p = re.sub(r"[^A-Z0-9 ]", " ", txt).split()
    k = 0
    while k < len(p) and p[k] in TIPOS_VIA:
        k += 1
    return k


def primeiro(mask, rotulo):
    d = lista[mask].head(1).copy()
    d["caso"] = rotulo
    return d


lista["n_tipos"] = lista["logradouro"].map(tipos_removidos)
lista["abrev"] = lista["logradouro"].map(lambda s: any(p in ABREVIACOES for p in re.sub(r"[^A-Z0-9 ]", " ", s).split()[1:]))
lista["pontuacao"] = lista["logradouro"].str.contains(r"[^A-Z0-9 ]")
linhas_chave = pd.concat([
    primeiro((lista["n_tipos"] == 1) & ~lista["abrev"], "tira o tipo de via"),
    primeiro((lista["n_tipos"] == 1) & lista["abrev"] & lista["logradouro"].str.contains(r"\bDR\b|\bPROF\b|\bPE\b|\bSTA\b|\bSTO\b"), "expande abreviação de título"),
    primeiro((lista["n_tipos"] == 1) & lista["logradouro"].str.contains(r"\bBRIG\b"), "expande BRIG → BRIGADEIRO"),
    primeiro(lista["pontuacao"] & lista["rua_key"].notna(), "pontuação vira espaço"),
    primeiro((lista["n_tipos"] >= 2) & lista["rua_key"].notna(), "dois tipos de via seguidos: os dois saem"),
    primeiro((lista["n_tipos"] == 0) & lista["rua_key"].notna(), "sem tipo de via: nada a tirar"),
    primeiro(lista["logradouro"].str.contains("VEDACAO"), "vedado → ∅"),
])
linhas_chave = linhas_chave[["caso", "logradouro", "rua_key", "n"]].rename(columns={"n": "ocorrências"})
linhas_chave["ocorrências"] = linhas_chave["ocorrências"].map(fm)
bloco_chave = caso(
    "chave_rua: do logradouro ao nome próprio da rua",
    "Fluxo: <code>sem_acento</code> → tira pontuação → remove tipos de via do começo (<code>R, AV, ALAMEDA, JARDIM…</code>) → expande abreviações de título (<code>BRIG, DR, PROF, STA…</code>). "
    f"O resultado é a chave de join com o ITBI. {pc(crimes['rua_key'].isna().mean())} dos crimes ficam sem chave.",
    f"{fm(crimes['rua_key'].notna().sum())} crimes com chave",
    render(linhas_chave, transposed=False, colcls={"rua_key": "novo"})
    + '<p class="nota"><b>Efeito colateral real:</b> <code>JARDIM</code>/<code>PARQUE</code>/<code>CONJUNTO</code> também estão na lista de tipos de via, então '
    "uma rua chamada <code>JARDIM AMERICA</code> perde o <i>Jardim</i> nos dois lados (crime e ITBI) — como a regra é a mesma, o join continua casando.</p>",
)
blocos_crime.append(bloco_chave)

secoes.append(passo(
    "s1", "1", "Crimes: planilha da SSP → <code>crimes_sp.parquet</code>",
    f"Cinco arquivos anuais, 2 abas semestrais cada. Cada bloco abaixo é um caso de borda real: a célula como veio da planilha (esquerda) e como ficou depois da limpeza (direita). "
    f"Resultado: <b>{fm(TOTAL_CRIMES)}</b> linhas × {crimes.shape[1] - 1} colunas.",
    [b for b in blocos_crime if b],
))

# ───────────────────────── 2 · itbi ─────────────────────────
arq = pd.DataFrame(RAW["itbi_arquivos"])
arq_tab = pd.DataFrame({
    "arquivo": arq["arquivo"], "abas de mês": arq["meses"].map(len),
    "abas ignoradas": arq["ignoradas"].map(lambda l: ", ".join(l)),
    "1ª célula da aba": arq["primeira_celula"].map(bruto),
    "cabeçalho no topo?": arq["cab_topo"].map(lambda v: "sim" if v else "NÃO"),
    "linhas mantidas": arq["mantidas"].map(fm), "descartadas": arq["descartadas"].map(fm),
})
desc_txt = []
for _, a in arq.iterrows():
    d = a["descartes"][0] if a["descartes"] else None
    desc_txt.append(H(f"{d['onde']}: {bruto(d['celula'])}") if d else "—")
arq_tab["1ª descartada"] = desc_txt
assert arq["mantidas"].sum() == TOTAL_ITBI, (arq["mantidas"].sum(), TOTAL_ITBI)

alv_itbi = [ALVOS["L22"], ALVOS["L24"], ALVOS["L26"]]
antes_i, fontes_i = df_bruto_itbi(alv_itbi)
depois_i = limpa_itbi_linhas(antes_i)
assert list(depois_i["sql"]) == alv_itbi
bloco_itbi_layout = caso(
    "Arquivos com layouts diferentes",
    "A planilha é lida <b>sem cabeçalho</b>, com os 28 nomes declarados à mão (<code>COLUNAS_ITBI</code>); só linhas cuja 1ª coluna é número entram. "
    "Isso resolve de uma vez: o arquivo de 2024 (cabeçalho ordenado para o fim da aba, porque texto ordena depois de número), rodapés e linhas vazias, e as abas de legenda.",
    f"{len(arq)} arquivos · {fm(arq['mantidas'].sum())} linhas",
    render(arq_tab, transposed=False)
    + f'<p class="nota">A soma da coluna «linhas mantidas» é exatamente as <b>{fm(TOTAL_ITBI)}</b> linhas de <code>itbi_sp.parquet</code>.</p>',
)
bloco_itbi_diff = caso(
    "Linha bruta (28 colunas) → linha tipada",
    "Uma guia de cada ano. Em <b>2024</b> o SQL vem como <i>float</i> (<code>100800297.0</code>): o zero à esquerda do código se perdeu na planilha, e <code>zfill(11)</code> o recupera — "
    "sem isso o <code>setor_fiscal</code> (3 primeiros dígitos) sairia errado. Cinco colunas são descartadas (vermelho); quatro são criadas (verde).",
    "28 → 27 colunas",
    diff_tabela(antes_i, depois_i, labels=[f"{f.split(' · ')[1]}" for f in fontes_i],
                ocultar=("vvr", "vvr_proporcional", "base_calculo", "valor_financiado", "situacao_sql", "area_terreno", "testada", "fracao_ideal", "uso_cod", "padrao_cod"))
    + '<p class="nota">Por brevidade, 10 colunas que passam sem alteração (vvr, vvr_proporcional, base_calculo, valor_financiado, situacao_sql, area_terreno, testada, fracao_ideal, uso_cod, padrao_cod) '
    "não aparecem; só mudam de <i>texto/float</i> para <i>float64</i>.</p>",
)


def ex_itbi(mask, n=2, cols=None, culpa=None):
    sel = itbi[mask].drop_duplicates("sql").sort_values("sql")
    d = sel.iloc[[0, len(sel) // 2]][
        cols or ["sql", "logradouro", "numero", "bairro_itbi", "valor_transacao", "proporcao_transmitida", "area_construida", "padrao_desc", "acc", "ano"]
    ].reset_index(drop=True)
    marks = {(i, culpa): "alvo" for i in range(len(d))} if culpa else {}
    return d, marks


c_prop = itbi["proporcao_transmitida"] > 100
c_num = itbi["numero"] == 99999
c_acc = itbi["acc"] > 2026
c_pad0 = itbi["padrao_desc"].astype(str) == "0"
c_torre = itbi["bairro_itbi"].fillna("").str.match(r"^(TORRE|BLOCO)\b")
c_ano = itbi["ano"] < 2022
lixos = []
for mask, titulo, regra, culpa in [
    (c_prop, "proporção transmitida acima de 100%", "Valor impossível (o máximo teórico é 100). Não é corrigido: o filtro <code>proporcao_transmitida == 100</code> do passo 6 elimina a linha.", "proporcao_transmitida"),
    (c_num, "número do imóvel = 99999", "Código de «sem número». <code>numero</code> não entra em nenhum join (só a rua), então não precisa de tratamento.", "numero"),
    (c_acc, "ano de construção no futuro", "Não pode ser corrigido. Em vez de descartar a linha, <code>idade</code> usa só <code>acc</code> entre 1800 e 2026; fora disso a idade fica ∅.", "acc"),
    (c_pad0, "padrão construtivo = '0'", "Texto sem significado. Cai no filtro «padrão residencial» (não começa com RESIDENCIAL).", "padrao_desc"),
    (c_torre, "campo Bairro com «TORRE 1» / «BLOCO B»", "O complemento do endereço foi digitado na coluna errada. A regra de vocabulário do passo 5.1 rejeita: só vale bairro que existe na base de crimes.", "bairro_itbi"),
    (c_ano, "transação antes de 2022", "Escrituras vêm desde 1994. Não há crime comparável antes de 2022; o filtro de ano do passo 6 remove.", "ano"),
]:
    d, m = ex_itbi(mask, culpa=culpa)
    lixos.append(caso(titulo, regra, f"{fm(mask.sum())} linhas ({pc(mask.mean(), 2)})", render(d, marks=m, transposed=False)))

secoes.append(passo(
    "s2", "2", "ITBI: planilha da Prefeitura → <code>itbi_sp.parquet</code>",
    f"Cinco arquivos (2022 a set/2026), uma aba por mês, 28 colunas. <b>{fm(TOTAL_ITBI)}</b> guias, só a capital. "
    "A célula laranja marca o valor problemático em cada caso.",
    [bloco_itbi_layout, bloco_itbi_diff] + lixos,
))

# ───────────────────────── 3 · pontes geográficas ─────────────────────────
dup_cod = setores_raw["cd_setor_fiscal"].astype(str).str.zfill(3).value_counts()
dup_cod = dup_cod[dup_cod > 1]
setores = setores_raw.assign(setor_fiscal=setores_raw["cd_setor_fiscal"].astype(str).str.zfill(3)).dissolve(by="setor_fiscal").reset_index()[["setor_fiscal", "geometry"]]
ov = gpd.overlay(setores, distritos, how="intersection", keep_geom_type=True)
ov["area_m2"] = ov.to_crs(31983).area
ov["parcela"] = ov["area_m2"] / ov.groupby("setor_fiscal")["area_m2"].transform("sum")
n_setor_trans = itbi["setor_fiscal"].astype(str).value_counts()
dom = ov.groupby("setor_fiscal")["parcela"].max()
dom_df = pd.DataFrame({"parcela": dom, "n": n_setor_trans}).dropna()
n_multi = int((ov.groupby("setor_fiscal")["nm_distrito_municipal"].nunique() > 1).sum())
media_pond = float((dom_df["parcela"] * dom_df["n"]).sum() / dom_df["n"].sum())
# o setor_para_distrito gravado deve coincidir com o recálculo
chk = ov.sort_values("area_m2").groupby("setor_fiscal")["nm_distrito_municipal"].last()
assert (chk.reindex(setor_dist["setor_fiscal"]).to_numpy() == setor_dist["distrito"].to_numpy()).mean() > 0.99

SETOR_LIMPO = dom_df[(dom_df["parcela"] > 0.99) & (dom_df["n"] > 5000)].sort_values("n").index[-1]
SETOR_CONF = dom_df[dom_df["n"] > 5000].sort_values("parcela").index[0]


def tab_setor(s):
    d = ov[ov["setor_fiscal"] == s].sort_values("area_m2", ascending=False)
    d = pd.DataFrame({
        "setor_fiscal": d["setor_fiscal"], "distrito": d["nm_distrito_municipal"],
        "área da interseção (m²)": d["area_m2"].round(0), "parcela do setor": d["parcela"].map(lambda x: pc(x, 1)),
    }).reset_index(drop=True)
    rc = {0: "novo"}
    return render(d, rowcls=rc, transposed=False)


buckets = [(0, .5), (.5, .7), (.7, .9), (.9, .99), (.99, 1.01)]
linhas_b = []
for lo, hi in buckets:
    m = (dom_df["parcela"] >= lo) & (dom_df["parcela"] < hi)
    rot = f"distrito dominante {'≥ 99%' if lo >= .99 else f'{int(lo * 100)}–{int(hi * 100)}%'} do setor"
    linhas_b.append((rot + f" · {int(m.sum())} setores", int(dom_df.loc[m, "n"].sum()), int(dom_df["n"].sum())))

bloco_geo1 = caso(
    "GeoSampa: duas camadas baixadas por WFS",
    "<code>distrito_municipal</code> (96) e <code>setor_fiscal</code> (feições). A geometria é simplificada a 20 m em projeção métrica (EPSG:31983) para o arquivo não pesar centenas de MB, e voltada a lat/lon.",
    f"{len(distritos)} distritos · {len(setores_raw)} feições de setor",
    f'<p class="nota">O mesmo código de setor aparece em mais de uma feição: <b>{", ".join(f"{k} (×{v})" for k, v in dup_cod.items())}</b>. '
    f"<code>dissolve(by='setor_fiscal')</code> funde as peças e resulta em <b>{len(setores)}</b> setores únicos (a chave do ITBI é única).</p>",
)
bloco_geo2 = caso(
    "Setor fiscal → distrito: por maior área de interseção",
    "<code>overlay</code> setor × distrito, soma a área de cada pedaço, fica o distrito com mais área. O ITBI só traz o setor (3 primeiros dígitos do SQL), então esta é a única ponte ITBI → distrito.",
    f"{len(setor_dist)} setores → {setor_dist['distrito'].nunique()} distritos",
    f'<p class="nota"><b>Setor {SETOR_LIMPO}</b> — caso limpo ({fm(dom_df.loc[SETOR_LIMPO, "n"])} guias):</p>{tab_setor(SETOR_LIMPO)}'
    f'<p class="nota"><b>Setor {SETOR_CONF}</b> — caso difícil ({fm(dom_df.loc[SETOR_CONF, "n"])} guias): o setor é maior que qualquer distrito e fica com o dominante, '
    f'mas só {pc(dom_df.loc[SETOR_CONF, "parcela"], 0)} da área está nele. Todas as guias do setor ganham o mesmo distrito.</p>{tab_setor(SETOR_CONF)}',
)
bloco_geo3 = caso(
    "Quão «exata» é essa ponte? Distribuição real",
    f"{n_multi} dos {len(setores)} setores cruzam mais de um distrito. Peso por número de guias: o distrito dominante cobre em média <b>{pc(media_pond, 0)}</b> da área do setor de uma guia típica. "
    "A afirmação «100% exato» do notebook 01 vale para <i>setor</i> (todo SQL tem setor), não para <i>distrito</i>. "
    "Consequência prática: o distrito só serve para <b>desambiguar rua</b> (camada <code>(rua_key, distrito)</code>), nunca como rótulo final.",
    f"{n_multi} setores multi-distrito",
    barra_tabela(linhas_b, cab=("faixa de dominância", "guias de ITBI", "", "% das guias")),
    "limite",
)

# sjoin crimes
corpo_sj = (
    barra_tabela([
        ("crimes com lat/lon (entram no sjoin)", n_pontos, TOTAL_CRIMES),
        ("→ caem em exatamente 1 distrito", n_pontos - n_fronteira - n_fora, TOTAL_CRIMES),
        ("→ caem em 2 polígonos (fronteira)", n_fronteira, TOTAL_CRIMES),
        ("→ caem fora de todos os distritos", n_fora, TOTAL_CRIMES),
        ("sem lat/lon (nem tentam o sjoin)", TOTAL_CRIMES - n_pontos, TOTAL_CRIMES),
    ], cab=("destino do crime", "linhas", "", "% dos crimes"))
)
if exemplo_fronteira is not None:
    ef = exemplo_fronteira.copy()
    ef["lon"], ef["lat"] = ef.geometry.x, ef.geometry.y
    ef = ef[["indice", "lat", "lon", "nm_distrito_municipal"]].rename(columns={"indice": "índice do crime", "nm_distrito_municipal": "distrito"})
    corpo_sj += (
        f'<p class="nota"><b>Fronteira:</b> o mesmo crime (índice {ef["índice do crime"].iloc[0]}) volta duas vezes do <code>sjoin</code> — as geometrias simplificadas se sobrepõem. '
        f"{fm(n_linhas_sjoin)} linhas de saída para {fm(n_pontos)} pontos; <code>drop_duplicates('indice')</code> fica com a primeira.</p>"
        + render(ef, transposed=False, rowcls={0: "novo", 1: "cai"})
    )
bloco_sj = caso(
    "Crime → distrito por ponto-em-polígono",
    "<code>gpd.sjoin(pontos, distritos, predicate='within')</code>. A coluna <code>distrito</code> existe só em memória; não entra em <code>crimes_sp.parquet</code>.",
    f"{pc(crimes['distrito'].notna().mean())} dos crimes ganham distrito", corpo_sj,
)
secoes.append(passo(
    "s3", "3", "Pontes geográficas: distrito para o ITBI e para o crime",
    "As duas bases não têm coordenada em comum: o ITBI só tem endereço e setor fiscal, o crime tem bairro, rua e (às vezes) lat/lon. "
    "O GeoSampa fornece a geometria que liga os dois a um terceiro nível, o <b>distrito</b>.",
    [bloco_geo1, bloco_geo2, bloco_geo3, bloco_sj],
))

# ───────────────────────── 4 · dicionário de ruas ─────────────────────────
cr_rua = crimes_com_rua_bairro[crimes_com_rua_bairro["rua_key"] == RUA]
cid_cont = cr_rua.groupby("bairro", observed=True).size().sort_values(ascending=False)
cid_cont = cid_cont[cid_cont > 0]
tab_cid = pd.DataFrame({"rua_key": RUA, "bairro (crimes desta rua)": cid_cont.index.astype(str), "ocorrências": cid_cont.values})
tab_cid["% da rua"] = (tab_cid["ocorrências"] / tab_cid["ocorrências"].sum()).map(pc)
tab_cid["modal?"] = ["← modal" if i == 0 else "" for i in range(len(tab_cid))]
d_rua = dic[dic["rua_key"] == RUA].copy()
d_rua["camada"] = np.where(d_rua["distrito"].isna(), "cidade-toda", "por distrito")
d_rua = d_rua.sort_values(["camada", "n_ocorrencias"], ascending=[False, False])[["camada", "rua_key", "distrito", "bairro", "n_ocorrencias", "lat", "lon"]]
d_rua = d_rua.head(8)
# confere o recálculo da camada cidade-toda
rec = cr_rua.groupby("bairro", observed=True).size().idxmax()
assert rec == dic[(dic["rua_key"] == RUA) & dic["distrito"].isna()]["bairro"].iloc[0]

bloco_dic1 = caso(
    f"Construção para uma rua: {RUA}",
    "<b>Cidade-toda:</b> agrupa por <code>(rua_key, bairro)</code>, conta, e o bairro com mais crimes vira o <i>modal</i>; lat/lon = <b>mediana</b> (uma coordenada errada não arrasta o centro). "
    "<b>Por distrito:</b> o mesmo cálculo, mas por <code>(rua_key, distrito)</code>, usando só os crimes que ganharam distrito no passo 3.",
    f"{fm(cid_cont.sum())} crimes na rua",
    render(tab_cid, transposed=False, rowcls={0: "novo"})
    + f'<p class="nota">Resultado em <code>dicionario_ruas.parquet</code> — {len(dic[dic["rua_key"] == RUA])} linhas para esta rua '
    "(1 cidade-toda + 1 por distrito; mostradas as primeiras 8):</p>" + render(d_rua, transposed=False, colcls={"bairro": "novo", "lat": "novo", "lon": "novo"}),
)

# Melo Alves
MA = "MELO ALVES"
ma_cr = crimes_com_rua_bairro[crimes_com_rua_bairro["rua_key"] == MA].groupby("bairro", observed=True).size()
ma_cr = ma_cr[ma_cr > 0]
ma_dic = dic[dic["rua_key"] == MA][["rua_key", "distrito", "bairro", "n_ocorrencias"]]
ma_it = loc[loc["rua_key"] == MA]
ma_res = ma_it.groupby(["fonte", "bairro"]).size().rename("guias").reset_index().sort_values("guias", ascending=False).head(6)
bloco_dic2 = caso(
    "Rua com poucos crimes: MELO ALVES (existe em dois bairros)",
    f"A rua existe nos Jardins e no Jardim Monte Verde, mas só <b>{int(ma_cr.sum())}</b> crimes foram registrados nela, todos num bairro. Pureza de 100% e mesmo assim errada para as guias dos Jardins. "
    f"Correção: o dicionário só vale com <b>≥ {MINIMO_OCORRENCIAS_RUA} ocorrências</b>; abaixo disso a rua é ignorada e o bairro vem do CEP ou do campo declarado.",
    f"{fm(len(ma_it))} guias na rua",
    f'<p class="nota">Crimes registrados em {MA}:</p>' + render(pd.DataFrame({"bairro": ma_cr.index.astype(str), "ocorrências": ma_cr.values}), transposed=False)
    + '<p class="nota">Linha no dicionário (n_ocorrencias &lt; 20 → <b>excluída</b> por <code>MINIMO_OCORRENCIAS_RUA</code> em <code>localiza()</code>):</p>' + render(ma_dic, transposed=False, colcls={"n_ocorrencias": "alvo"})
    + '<p class="nota">O que as guias de ' + MA + ' recebem, e de qual fonte:</p>' + render(ma_res.rename(columns={"fonte": "fonte que decidiu"}), transposed=False),
)

# pureza
RUAS_AMB = ["RAPOSO TAVARES", "PRESIDENTE WILSON", "BRIGADEIRO LUIS ANTONIO", "MELO ALVES"]
linhas_p = []
for r in RUAS_AMB:
    c = crimes_com_rua_bairro.loc[crimes_com_rua_bairro["rua_key"] == r, "bairro"].astype(str).value_counts()
    cd = crimes_com_rua_bairro[(crimes_com_rua_bairro["rua_key"] == r) & crimes_com_rua_bairro["distrito"].notna()]
    pur_d = cd.groupby("distrito").apply(lambda g: g["bairro"].astype(str).value_counts().iloc[0] / len(g), include_groups=False)
    linhas_p.append({
        "rua_key": r, "ocorrências": fm(c.sum()), "bairros distintos": len(c),
        "bairro modal (cidade-toda)": c.index[0], "pureza cidade-toda": pc(c.iloc[0] / c.sum()),
        "distritos com crime": len(pur_d), "pureza média por distrito": pc(pur_d.mean()),
    })
bloco_dic3 = caso(
    "Via longa: a mesma rua cruza vários bairros",
    "<i>Pureza</i> = fração dos crimes da rua que cai no bairro modal. Rua que atravessa a cidade tem pureza baixa na camada cidade-toda; "
    "dentro de um distrito ela sobe, e é por isso que a chave de preferência é <code>(rua_key, distrito)</code>.",
    f"{fm(dic['distrito'].isna().sum())} ruas cidade-toda · {fm(dic['distrito'].notna().sum())} pares rua×distrito",
    render(pd.DataFrame(linhas_p), transposed=False, colcls={"pureza cidade-toda": "mod", "pureza média por distrito": "novo"}),
)
secoes.append(passo(
    "s4", "4", "<code>dicionario_ruas.parquet</code>: de crime a «rua → bairro»",
    f"{fm(len(crimes_com_rua_bairro))} crimes com rua e bairro viram {fm(len(dic))} linhas: uma camada cidade-toda (fallback) e uma camada por distrito.",
    [bloco_dic1, bloco_dic2, bloco_dic3],
))

# ───────────────────────── 5 · localiza ─────────────────────────
N = len(loc)
# 5.1 normalização
distintos = loc.drop_duplicates("bairro_itbi").dropna(subset=["bairro_itbi"])
cnt_b = loc["bairro_itbi"].value_counts()
dist_b = pd.DataFrame({"bairro_itbi": distintos["bairro_itbi"], "norm": distintos["bairro_norm"], "ok": distintos["bairro_norm"].isin(VOCAB)})
dist_b["guias"] = dist_b["bairro_itbi"].map(cnt_b)


def pick_b(mask, rot):
    d = dist_b[mask].sort_values("guias", ascending=False).head(1).copy()
    d["caso"] = rot
    return d


cases_b = pd.concat([
    pick_b((dist_b["bairro_itbi"] == dist_b["norm"]) & dist_b["ok"], "já limpo e existe no vocabulário"),
    pick_b((dist_b["bairro_itbi"] != dist_b["norm"]) & dist_b["ok"], "abreviação expandida, passa"),
    pick_b((dist_b["bairro_itbi"].str.match(r"^(VL|JD|PQ|CJ|CID|STO|STA|CH)\b")) & ~dist_b["ok"] & (dist_b["bairro_itbi"] != dist_b["norm"]), "expande, mas o bairro não existe nos crimes"),
    pick_b(dist_b["bairro_itbi"].str.match(r"^(TORRE|BLOCO)\b"), "complemento digitado no campo Bairro"),
    pick_b(dist_b["bairro_itbi"].str.contains(r"[,\d]") & ~dist_b["bairro_itbi"].str.match(r"^(TORRE|BLOCO)\b"), "endereço no campo Bairro"),
    pick_b(dist_b["bairro_itbi"].str.match(r"^N SRA") & ~dist_b["ok"], "abreviação que a tabela não cobre"),
])
cases_b["resultado"] = np.where(cases_b["ok"], "VALE → " + cases_b["norm"].astype(str), "rejeitado → ∅")
cases_b["guias"] = cases_b["guias"].map(fm)
cases_b = cases_b[["caso", "bairro_itbi", "norm", "guias", "resultado"]].rename(columns={"norm": "normalizado"})
vazias_b = int(loc["bairro_itbi"].isna().sum())
n_dec = int(loc["bairro_declarado"].notna().sum())
bloco_51 = caso(
    "5.1 · Bairro declarado na guia, validado pelo vocabulário dos crimes",
    f"<code>normaliza_bairro</code> (sem acento, pontuação → espaço, <code>VL→VILA, JD→JARDIM, STO→SANTO…</code>) e só vale se o resultado existir entre os <b>{fm(len(VOCAB))}</b> bairros da base de crimes. "
    f"{pc(loc['bairro_itbi'].notna().mean())} das guias trazem algum texto ({fm(vazias_b)} vazias); {pc(n_dec / N)} passam na validação. Sem esse filtro, «TORRE 1» viraria bairro.",
    f"{fm(loc['bairro_itbi'].nunique())} textos distintos → {fm(int(dist_b['ok'].sum()))} válidos",
    render(cases_b, transposed=False, colcls={"resultado": "mod"}),
)

# 5.2 CEP
vot = loc.dropna(subset=["bairro_declarado"]).groupby(["cep_txt", "bairro_declarado"]).size().rename("votos").reset_index()
agg = vot.groupby("cep_txt")["votos"].agg(["sum", "max"])
agg["parcela"] = agg["max"] / agg["sum"]
disputa = agg[(agg["sum"] >= 20)].sort_values("parcela")
CEP_EX = disputa.index[0]
vot_ex = vot[vot["cep_txt"] == CEP_EX].sort_values("votos", ascending=False)
vot_ex["resultado"] = ["← vence" if i == 0 else "" for i in range(len(vot_ex))]
guias_cep = loc[loc["cep_txt"] == CEP_EX]
sem_decl = guias_cep[guias_cep["bairro_declarado"].isna()]
ex_sem = sem_decl[["sql", "logradouro", "numero", "bairro_itbi", "cep_txt", "bairro_declarado", "bairro_cep", "bairro"]].head(3)
disc = loc[loc["bairro_declarado"].notna() & loc["bairro_cep"].notna() & (loc["bairro_declarado"] != loc["bairro_cep"])]
disc_ex = disc.sort_values("sql").iloc[[0, len(disc) // 2]][["sql", "logradouro", "numero", "cep_txt", "bairro_declarado", "bairro_cep", "bairro"]]
cep_nulo = loc["cep"].isna()
bloco_52 = caso(
    "5.2 · Propagação por CEP (voto)",
    "Para cada CEP, o bairro <i>declarado e válido</i> mais frequente entre as guias daquele CEP vence (CEP ≈ quadra). Ele preenche as guias do mesmo CEP que não declararam bairro. "
    "Quem declarou um bairro válido <b>mantém o seu</b> (prioridade 1).",
    f"{fm(int(loc['bairro_cep'].notna().sum()))} guias com voto de CEP ({pc(loc['bairro_cep'].notna().mean())})",
    f'<p class="nota"><b>Eleição disputada, CEP {CEP_EX}</b> — {int(agg.loc[CEP_EX, "sum"])} guias com bairro válido, o vencedor tem só {pc(agg.loc[CEP_EX, "parcela"], 0)} dos votos:</p>'
    + render(vot_ex.rename(columns={"bairro_declarado": "bairro declarado"}), transposed=False)
    + '<p class="nota">Guias deste CEP sem bairro válido, e o que recebem:</p>'
    + render(ex_sem.rename(columns={"bairro_declarado": "declarado (válido)", "bairro_cep": "voto do CEP", "bairro": "bairro final"}), transposed=False, colcls={"bairro_cep": "novo", "voto do CEP": "novo", "bairro final": "novo"})
    + f'<p class="nota"><b>Quando o declarado e o CEP discordam</b> ({fm(len(disc))} guias): o declarado vence.</p>'
    + render(disc_ex.rename(columns={"bairro_declarado": "declarado (válido)", "bairro_cep": "voto do CEP", "bairro": "bairro final"}), transposed=False, colcls={"declarado (válido)": "novo"}),
)

# 5.3 rua
usado3 = loc[loc["fonte"].isin(["3a · rua + distrito", "3b · rua cidade-toda"])]
ex3a = loc[loc["fonte"] == "3a · rua + distrito"].sort_values("sql").iloc[[0, len(loc[loc["fonte"] == "3a · rua + distrito"]) // 2]]
ex3b = loc[loc["fonte"] == "3b · rua cidade-toda"].sort_values("sql").iloc[[0, len(loc[loc["fonte"] == "3b · rua cidade-toda"]) // 2]]
cols3 = ["sql", "logradouro", "rua_key", "distrito", "bairro_local", "bairro_cidade", "bairro"]
ex3 = pd.concat([ex3a, ex3b])[cols3]
sem = loc[loc["fonte"] == "sem bairro"]
motivos = pd.Series(np.select(
    [sem["rua_key"].isna(), ~sem["rua_key"].isin(set(dic["rua_key"])), True],
    ["sem rua_key (logradouro vazio)", "rua nunca aparece nos crimes", f"rua existe, mas < {MINIMO_OCORRENCIAS_RUA} ocorrências na cidade e no distrito"], "?"), index=sem.index)
tab_mot = motivos.value_counts().rename("guias").reset_index().rename(columns={"index": "motivo", 0: "motivo"})
tab_mot.columns = ["motivo", "guias"]
sem_ex = sem.assign(motivo=motivos).sort_values("sql").groupby("motivo").head(1)[["sql", "logradouro", "rua_key", "cep_txt", "bairro_itbi", "motivo"]]
bloco_53 = caso(
    "5.3 · Dicionário de ruas (último recurso)",
    f"Só entra o que sobrou. Chave <code>(rua_key, distrito)</code> primeiro; se a rua não tem crimes suficientes <i>naquele distrito</i> (&lt; {MINIMO_OCORRENCIAS_RUA}), cai para a rua cidade-toda. "
    "Quando as duas têm resposta, a do distrito vence (<code>bairro_local.fillna(bairro_cidade)</code>).",
    f"{fm(len(usado3))} guias decididas aqui ({pc(len(usado3) / N)})",
    render(ex3.rename(columns={"bairro_local": "bairro (rua+distrito)", "bairro_cidade": "bairro (cidade-toda)", "bairro": "bairro final"}), transposed=False, colcls={"bairro final": "novo"})
    + f'<p class="nota"><b>Guias sem bairro por nenhuma das três fontes</b> ({fm(len(sem))} = {pc(len(sem) / N)}): ficam ∅ e saem no filtro «bairro identificado».</p>'
    + render(tab_mot, transposed=False) + render(sem_ex, transposed=False),
)

# 5.4 coordenada
coord_ok = loc["bairro"].notna() & loc["rua_concorda"]
coord_centro = loc["bairro"].notna() & ~loc["rua_concorda"]
cols4 = ["sql", "logradouro", "bairro", "bairro_rua", "rua_concorda", "lat_rua", "lon_rua", "lat_centro", "lon_centro", "lat", "lon"]
e_ok = loc[coord_ok & (loc["fonte"] == "3a · rua + distrito")].sort_values("sql").iloc[[0]]
e_dif = loc[coord_centro & loc["bairro_rua"].notna() & (loc["fonte"] == "2 · CEP")].sort_values("sql").iloc[[0]]
e_sem = loc[coord_centro & loc["bairro_rua"].isna() & (loc["fonte"] == "2 · CEP")].sort_values("sql").iloc[[0]]
ex4 = pd.concat([e_ok, e_dif, e_sem])[cols4].reset_index(drop=True)
ex4["rua_concorda"] = ex4["rua_concorda"].map({True: "sim", False: "não"})
bloco_54 = caso(
    "5.4 · Coordenada: a da rua só vale se concordar com o bairro",
    "A coordenada da rua (mediana dos crimes dela) só é aceita se o bairro modal da rua <b>for igual ao bairro final</b>. Se a rua aponta para outro bairro (porque o CEP/declarado mandou), usar a coordenada dela colocaria a guia no lugar errado: "
    "usa-se o <b>centro do bairro final</b> (mediana de lat/lon dos crimes geolocalizados dele).",
    f"{pc(coord_ok.sum() / N)} usam a rua · {pc(coord_centro.sum() / N)} usam o centro do bairro",
    render(ex4, colcls={"lat": "novo", "lon": "novo"}, transposed=True, labels=["rua concorda", "rua discorda", "rua nem está no dicionário"]),
)

casc = [
    ("1 · bairro declarado na guia", int((loc["fonte"] == "1 · bairro declarado").sum()), N),
    ("2 · + voto do CEP", int((loc["fonte"] == "2 · CEP").sum()), N),
    ("3a · + rua com distrito", int((loc["fonte"] == "3a · rua + distrito").sum()), N),
    ("3b · + rua cidade-toda", int((loc["fonte"] == "3b · rua cidade-toda").sum()), N),
    ("sem bairro", int((loc["fonte"] == "sem bairro").sum()), N),
]
assert sum(c[1] for c in casc) == N
acum, ac = [], 0
for nome, q, _ in casc[:-1]:
    ac += q
    acum.append(ac / N)
bloco_casc = caso(
    "Cascata: quem decidiu o bairro de cada guia (contagem exclusiva)",
    f"Cada guia aparece em <b>uma</b> linha: a primeira fonte que respondeu. Cobertura acumulada: {pc(acum[0])} → {pc(acum[1])} (com CEP) → <b>{pc(acum[3])}</b> no total.",
    f"{fm(N)} guias",
    barra_tabela(casc, cab=("fonte que decidiu", "guias", "", "% das guias")),
)
# joins do localiza
njoin = pd.DataFrame([
    ("ITBI ⨝ setor_para_distrito", "setor_fiscal", "N → 1", N, N, int(loc["distrito"].notna().sum())),
    ("ITBI ⨝ voto por CEP", "cep_txt", "N → 1", N, N, int(loc["bairro_cep"].notna().sum())),
    ("ITBI ⨝ dicionário (rua, distrito)", "rua_key + distrito", "N → 1", N, N, int(loc["bairro_local"].notna().sum())),
    ("ITBI ⨝ dicionário (cidade-toda)", "rua_key", "N → 1", N, N, int(loc["bairro_cidade"].notna().sum())),
], columns=["merge", "chave", "cardinalidade", "linhas antes", "linhas depois", "linhas que casaram"])
njoin["% casou"] = (njoin["linhas que casaram"] / njoin["linhas depois"]).map(pc)
for c in ("linhas antes", "linhas depois", "linhas que casaram"):
    njoin[c] = njoin[c].map(fm)
bloco_joins = caso(
    "Os quatro merges de localiza()", "Todos são <i>left joins N→1</i>: o número de linhas do ITBI não muda (assert no script); só aparecem colunas novas com ∅ onde não casou.",
    "linhas preservadas", render(njoin, transposed=False), "info",
)
ex_cep = loc[loc["sql"].isin([ALVOS["A"], ALVOS["B"], ALVOS["C"]])][["sql", "cep", "cep_txt"]].reset_index(drop=True)
bloco_cep_zero = caso(
    "CEP numérico perde o zero à esquerda",
    "A planilha guarda o CEP como número: <code>03010-000</code> vira <code>3010000</code>. <code>cep.astype('Int64').astype(str).str.zfill(8)</code> recompõe os 8 dígitos antes de agrupar, para o CEP ter a mesma chave em todas as guias.",
    f"{fm((loc['cep'] < 10_000_000).sum())} guias com CEP &lt; 10 milhões",
    render(ex_cep, transposed=False, colcls={"cep_txt": "novo"}),
)
secoes.append(passo(
    "s5", "5", "<code>localiza()</code>: dando bairro e coordenada a cada guia",
    "O ITBI só declara bairro em parte das guias, e de forma suja. Três fontes em ordem; a primeira que responde decide. "
    "Cada bloco mostra linhas reais daquela fonte.",
    [bloco_51, bloco_cep_zero, bloco_52, bloco_53, bloco_54, bloco_casc, bloco_joins],
))

# ───────────────────────── 6 · filtros ─────────────────────────
linhas_f = [(funil[0][0], funil[0][1], funil[0][1])]
for (n_, q) in funil[1:]:
    linhas_f.append((n_, q, funil[0][1]))
blocos_f = [caso("Funil", "Cada filtro vale para as linhas que sobraram do anterior. Contagens recalculadas agora a partir de <code>itbi_sp.parquet</code> e conferidas contra <code>base_modelagem.parquet</code>.", f"{fm(funil[0][1])} → {fm(funil[-1][1])}", barra_tabela(linhas_f, cab=("depois do filtro", "guias", "", "% do total")), "info")]

TEXTOS_FILTRO = {
    "transação entre 2022 e 2026": "Sem crime comparável antes de 2022.",
    "natureza = compra e venda": "Dação em pagamento, permuta, leilão, integralização de capital etc. não têm preço de mercado.",
    "100% do imóvel transmitido": "Abaixo de 100%, a guia é fração ideal; a área construída pertence ao prédio inteiro. Acima de 100% é valor inválido.",
    "padrão residencial": "Terreno, comercial, industrial e templo saem. É a hipótese: preço de moradia.",
    "exclui vaga de garagem": "Vaga vendida como unidade autônoma tem área pequena e valor pequeno: infla/destorce o R$/m².",
    "área construída e valor > 0": "Protege a divisão por zero do R$/m². Os terrenos (área 0) já saíram no filtro de padrão residencial, então sobra quase nada aqui.",
    "bairro identificado": "Sem bairro não há como colar o perfil de crime.",
    "mês coberto pelo IPCA": "Mês sem índice ainda não é estimado, é descartado.",
}
for nome, _, cols in FILTROS:
    rem = removidas[nome]
    n_in = int(funil[[f[0] for f in funil].index(nome) - 1][1])
    if len(rem) == 0:
        corpo = '<p class="nota">Nenhuma linha removida neste passo.</p>'
    else:
        c0 = cols[0]
        serie0 = rem[c0].astype("Int64") if c0 == "ano" else rem[c0]
        top = serie0.astype(str).value_counts().head(5).rename_axis(c0).reset_index(name="guias removidas")
        top["guias removidas"] = top["guias removidas"].map(fm)
        rem_u = rem.drop_duplicates("sql").sort_values("sql")
        amostra = rem_u.iloc[sorted(set(np.linspace(0, len(rem_u) - 1, 3).astype(int)))][["sql", "logradouro", "numero"] + [c for c in cols if c not in ("logradouro", "numero")]].reset_index(drop=True)
        marks = {(i, c0): "alvo" for i in range(len(amostra))}
        corpo = (
            f'<p class="nota">Motivos mais comuns (coluna <code>{c0}</code>):</p>' + render(top, transposed=False)
            + '<p class="nota">Três guias reais removidas aqui (a célula laranja é o motivo):</p>' + render(amostra, marks=marks, transposed=False)
        )
    blocos_f.append(caso(nome, TEXTOS_FILTRO[nome], f"−{fm(len(rem))} ({pc(len(rem) / n_in, 2)} das {fm(n_in)})", corpo))

# percentis
mm = fora_pct.sort_values("preco_m2")
amostra_p = pd.concat([mm.head(2), mm.tail(2)])[["sql", "logradouro", "valor_transacao", "area_construida", "valor_real", "preco_m2"]].reset_index(drop=True)
blocos_f.append(caso(
    "R$/m² real fora dos percentis 1 e 99",
    f"Corte calculado <b>depois</b> da deflação: p1 = <b>{brl(P1)}</b>/m², p99 = <b>{brl(P99)}</b>/m². O corte inferior é baixo: só pega valores simbólicos (R$ 0,01 declarado); o superior pega erros de digitação de área ou valor.",
    f"−{fm(len(fora_pct))} ({pc(len(fora_pct) / n_antes_pct, 2)})",
    render(amostra_p, marks={(i, "preco_m2"): "alvo" for i in range(len(amostra_p))}, transposed=False)
    + f'<p class="nota">Piso do corte: {brl(P1)}; teto: {brl(P99)}. Na base final, mínimo {brl(base["preco_m2"].min())} e máximo {brl(base["preco_m2"].max())}.</p>',
))
# idade
acc_ruim = base[base["acc"].notna() & ~base["acc"].between(1800, 2026)]
acc_nan = int(base["acc"].isna().sum())
exemplo_idade = base[base["acc"].between(1800, 2026)].sort_values("sql").iloc[[0]]
blocos_f.append(caso(
    "Idade do imóvel: ano de construção inválido não derruba a linha",
    "<code>idade = ano − acc</code>, mas só com <code>acc</code> entre 1800 e 2026. Fora disso a idade é ∅ (a linha continua na base; só não participa da mediana de idade do bairro).",
    f"{len(acc_ruim)} acc inválidos · {fm(acc_nan)} sem acc",
    render(pd.concat([exemplo_idade, acc_ruim.head(1)] if len(acc_ruim) else [exemplo_idade])[["sql", "ano", "acc", "idade"]], transposed=False, colcls={"idade": "novo"}),
))
secoes.append(passo(
    "s6", "6", "Filtros: de guia de ITBI a imóvel residencial comparável",
    "Oito filtros em sequência, depois o corte de percentis. Para cada um: quantas linhas saem, por quê, e três exemplos reais removidos.",
    blocos_f,
))
assert len(base) == funil[-1][1]

# ───────────────────────── 7 · IPCA ─────────────────────────
rA = modelagem[modelagem["sql"] == SQL_A].iloc[0]
meses_ex = sorted({rA["mes"], ipca["mes"].iloc[-2], MES_REF})
ip_ex = ipca[ipca["mes"].isin(meses_ex)][["periodo", "indice", "variacao_pct", "mes", "fator"]].reset_index(drop=True)
ip_raw = ipca_csv[ipca_csv["periodo"].isin(ip_ex["periodo"])][["periodo", "indice", "variacao_pct"]].reset_index(drop=True)
idx_ref = ipca["indice"].iloc[-1]
idx_mes = float(ipca.loc[ipca["mes"] == rA["mes"], "indice"].iloc[0])
sem_ipca = removidas["mês coberto pelo IPCA"]
anual = ipca.assign(ano=ipca["mes"].str[:4]).groupby("ano").agg(
    meses=("variacao_pct", "size"), indice_dez=("indice", "last"),
    por_encadeamento=("variacao_pct", lambda v: ((1 + v / 100).prod() - 1) * 100),
)
anual["por_numero_indice"] = (anual["indice_dez"] / anual["indice_dez"].shift() - 1) * 100
anual["diferença"] = anual["por_encadeamento"] - anual["por_numero_indice"]
anual_t = anual.reset_index()
tab_anual = pd.DataFrame({
    "ano": anual_t["ano"], "meses": anual_t["meses"],
    "IPCA por encadeamento": anual_t["por_encadeamento"].map(lambda v: f"{v:.2f}%".replace(".", ",")),
    "IPCA por número-índice": anual_t["por_numero_indice"].map(lambda v: "—" if pd.isna(v) else f"{v:.2f}%".replace(".", ",")),
    "diferença": anual_t["diferença"].map(lambda v: "—" if pd.isna(v) else ("0,000 pp" if abs(v) < 0.0005 else f"{v:+.3f} pp".replace(".", ","))),
})
bloco_ipca1 = caso(
    "Linhas do IPCA usadas na guia A",
    f"Fator = índice do mês de referência ({MES_REF}) ÷ índice do mês da transação. Não se encadeia as variações mensais (acumularia erro de arredondamento).",
    f"{len(ipca)} meses na série",
    diff_tabela(ip_raw, ip_ex, labels=[f"{m}" for m in ip_ex["mes"]], antes_bruto=False)
    + f'<p class="nota">Guia A: transação em <b>{rA["mes"]}</b> → fator = {idx_ref:.2f} ÷ {idx_mes:.2f} = <b>{rA["fator"]:.6f}</b>.</p>',
)
bloco_ipca2 = caso(
    "A conta, aberta, para a guia A",
    "Merge <code>base ⨝ ipca[mes, fator]</code> (N → 1, chave <code>mes</code>), depois as três colunas de preço.",
    f"sql {SQL_A}",
    render(pd.DataFrame([
        ("valor_transacao", f"{rA['valor_transacao']:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."), "nominal, como declarado"),
        ("× fator", f"{rA['fator']:.6f}", f"índice {MES_REF} ÷ índice {rA['mes']}"),
        ("= valor_real", f"{rA['valor_real']:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."), f"R$ de {MES_REF}"),
        ("÷ area_construida", f"{rA['area_construida']:.0f} m²", ""),
        ("= preco_m2_nominal", f"{rA['preco_m2_nominal']:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."), "sem deflacionar"),
        ("= preco_m2", f"{rA['preco_m2']:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."), "valor usado em todo o resto"),
    ], columns=["passo", "valor", "significado"]), transposed=False, rowcls={2: "novo", 5: "novo"}),
)
bloco_ipca3 = caso(
    "Mês sem índice no IPCA",
    f"A série vai até {MES_REF}; a planilha de ITBI vai até {itbi['mes'].max()}. Guias desses meses não têm fator e são descartadas, não estimadas.",
    f"−{fm(len(sem_ipca))} guias",
    render(sem_ipca.sort_values("sql").head(2)[["sql", "logradouro", "data_transacao", "mes", "valor_transacao"]], transposed=False, marks={(0, "mes"): "alvo", (1, "mes"): "alvo"}) if len(sem_ipca) else "",
)
bloco_ipca4 = caso(
    "Checagem da série: IPCA anual por dois caminhos independentes",
    "Razão de índices × produto das variações mensais. Se concordam, a série foi lida certo.",
    "diferença máxima " + f"{anual['diferença'].abs().max():.3f} pp".replace(".", ","),
    render(tab_anual, transposed=False), "info",
)
secoes.append(passo(
    "s7", "7", "IPCA: deflação para reais constantes",
    f"Preço em reais de {MES_REF}. Fonte: API do IBGE (tabela 1737), variáveis 2266 (número-índice) e 63 (variação mensal).",
    [bloco_ipca1, bloco_ipca2, bloco_ipca3, bloco_ipca4],
))

# ───────────────────────── 8 · grupos de crime ─────────────────────────
nat = crimes["natureza"].astype(str).value_counts().rename_axis("natureza").reset_index(name="n")
nat["trânsito/culposo?"] = nat["natureza"].map(lambda s: "ACIDENTE DE TRANSITO" in s or "CULPOS" in s)
nat["furto?"] = nat["natureza"].str.startswith("FURTO")
nat["violento?"] = nat["natureza"].map(lambda s: any(t in s for t in VIOLENTOS))
nat["grupo final"] = nat["natureza"].map(grupo_crime)
nat["conflito de regras"] = (nat[["trânsito/culposo?", "furto?", "violento?"]].sum(axis=1) > 1)
nat_t = nat.copy()
for c in ["trânsito/culposo?", "furto?", "violento?", "conflito de regras"]:
    nat_t[c] = nat_t[c].map({True: "sim", False: ""})
nat_t["n"] = nat_t["n"].map(fm)
nat_t = nat_t.sort_values(["grupo final", "natureza"])
marks_n = {}
for i, (_, r) in enumerate(nat_t.reset_index(drop=True).iterrows()):
    if r["conflito de regras"] == "sim":
        marks_n[(i, "conflito de regras")] = "alvo"
grupos_tot = crimes["grupo"].value_counts().rename_axis("grupo").reset_index(name="ocorrências")
grupos_tot["%"] = (grupos_tot["ocorrências"] / len(crimes)).map(pc)
grupos_tot["ocorrências"] = grupos_tot["ocorrências"].map(fm)
bloco_g1 = caso(
    f"{len(nat)} naturezas → 4 grupos, por regra em ordem",
    "A ordem é a regra: trânsito/culposo primeiro, depois furto, depois violento, o resto «outros». Sem isso, <code>HOMICIDIO CULPOSO POR ACIDENTE DE TRANSITO</code> contaria como violência por conter «HOMICIDIO». "
    "Colunas «sim» mostram quais regras <i>casariam</i> o texto; a primeira da lista decide.",
    f"{int(nat['conflito de regras'].sum())} naturezas casam com mais de uma regra",
    render(nat_t.reset_index(drop=True), marks=marks_n, colcls={"grupo final": "novo"}, transposed=False),
)
_nomes = set(nat["natureza"])
dupl_nat = [n for n in _nomes if "  " in n and n.replace("  ", " - ") in _nomes] + [n.replace("  ", " - ") for n in _nomes if "  " in n and n.replace("  ", " - ") in _nomes]
bloco_g2 = caso(
    "Total por grupo", "O grupo é o que vira coluna em <code>bairros.parquet</code> (média anual 2022–2025).", f"{fm(len(crimes))} crimes",
    render(grupos_tot, transposed=False, colcls={"grupo": "novo"}), "info",
)
extra_g = ""
if dupl_nat:
    extra_g = caso(
        "Mesma natureza, grafia diferente",
        "Duas grafias da mesma natureza coexistem (traço presente ou ausente). As regras do grupo usam <code>in</code>/<code>startswith</code> sobre palavras-chave, então ambas caem no mesmo grupo sem normalização extra.",
        f"{len(dupl_nat)} nomes",
        render(nat[nat["natureza"].isin(dupl_nat)][["natureza", "n", "grupo final"]].assign(n=lambda d: d["n"].map(fm)), transposed=False),
    )
secoes.append(passo(
    "s8", "8", "Grupos de crime",
    "Furto mede fluxo de gente; roubo, com vítima, mede risco. Por isso 25 naturezas viram 4 grupos antes de agregar por bairro.",
    [bloco_g1, bloco_g2] + ([extra_g] if extra_g else []),
))

# ───────────────────────── 9 · agregação por bairro ─────────────────────────
BX = "MOOCA"
cx = por_bairro_grupo_ano.loc[BX].to_frame().T if BX in por_bairro_grupo_ano.index else None
cx = crimes[(crimes["bairro"].astype(str) == BX) & crimes["ano"].isin(ANOS_COMPLETOS)].groupby(["grupo", "ano"]).size().unstack(fill_value=0)
cx.columns = [str(int(c)) for c in cx.columns]
cx["média/ano"] = cx.sum(axis=1) / len(ANOS_COMPLETOS)
cx = cx.reset_index()
tot_linha = {"grupo": "TOTAL"}
for c in cx.columns[1:]:
    tot_linha[c] = cx[c].sum()
cx = pd.concat([cx, pd.DataFrame([tot_linha])], ignore_index=True)
bx_b = base[base["bairro"] == BX]
rb = bairros.loc[BX]
assert abs(bx_b["preco_m2"].median() - rb["preco_m2_mediano"]) < 1e-6
assert abs(cx["média/ano"].iloc[-1] - rb["crimes_total"]) < 1e-6
der = pd.DataFrame([
    ("preco_m2_mediano", f"mediana de preco_m2 das {fm(len(bx_b))} guias do bairro", brl(rb["preco_m2_mediano"])),
    ("n_transacoes", "nº de guias do bairro na base final", fm(rb["n_transacoes"])),
    ("crimes_total", "soma dos 4 grupos ÷ 4 anos", num(rb["crimes_total"])),
    ("violentos", "grupo «violento» ÷ 4 anos", num(rb["violentos"])),
    ("share_violento", f"violentos ÷ crimes_total = {num(rb['violentos'])} ÷ {num(rb['crimes_total'])}", num(rb["share_violento"], 4)),
    ("violentos_por_1k", f"violentos ÷ n_transacoes × 1000 = {num(rb['violentos'])} ÷ {fm(rb['n_transacoes'])} × 1000", num(rb["violentos_por_1k"])),
    ("log_violentos", "log10(1 + violentos): evita −∞", num(rb["log_violentos"], 4)),
], columns=["coluna", "como foi calculada", f"valor para {BX}"])
bloco_b1 = caso(
    f"Um bairro aberto: {BX}",
    "Crimes: só anos completos (2022–2025; 2026 é parcial), contados por grupo e divididos por 4. Preço: mediana (resiste a outlier) do R$/m² real. As duas metades se juntam por <b>nome do bairro</b> (<code>join</code> inner).",
    f"{fm(len(bx_b))} guias · {fm(int(cx.iloc[-1, 1:5].sum()))} crimes em 4 anos",
    '<p class="nota">Contagem de crimes por grupo e ano:</p>' + render(cx, transposed=False, rowcls={len(cx) - 1: "novo"})
    + '<p class="nota">Colunas derivadas:</p>' + render(der, transposed=False),
)

# bairros descartados
lado_preco = set(preco_bairro.index)
lado_crime = set(por_bairro_grupo.index)
so_preco = lado_preco - lado_crime
juntos = bair
fracos_t = juntos[juntos["n_transacoes"] < 30]
fracos_c = juntos[(juntos["n_transacoes"] >= 30) & (juntos["crimes_total"] < 30)]
ex_desc = pd.concat([
    fracos_t.sort_values("n_transacoes", ascending=False).head(3).assign(motivo="< 30 guias"),
    fracos_c.sort_values("crimes_total", ascending=False).head(3).assign(motivo="< 30 crimes/ano"),
])[["n_transacoes", "crimes_total", "motivo"]].reset_index().rename(columns={"index": "bairro"})
ex_desc["bairro"] = ex_desc["bairro"].astype(str)
ex_desc["n_transacoes"] = ex_desc["n_transacoes"].map(fm)
ex_desc["crimes_total"] = ex_desc["crimes_total"].map(num)
funil_b = [
    ("bairros com ao menos 1 guia na base filtrada", len(lado_preco), len(lado_preco)),
    ("… que também aparecem nos crimes (2022–2025)", len(juntos), len(lado_preco)),
    ("… com ≥ 30 guias", int((juntos["n_transacoes"] >= 30).sum()), len(lado_preco)),
    ("… e ≥ 30 crimes/ano  → bairros.parquet", len(bair_aprovados), len(lado_preco)),
]
bloco_b2 = caso(
    "Bairros descartados por falta de volume",
    "Mediana e taxa de bairro com poucas guias ou poucos crimes são ruído. O corte é <b>30 guias e 30 crimes/ano</b> (média). Um bairro que só aparece de um lado do join sai no <code>inner</code>.",
    f"{len(lado_preco) - len(bair_aprovados)} bairros fora · {len(bair_aprovados)} ficam",
    barra_tabela(funil_b, cab=("etapa", "bairros", "", "%"))
    + '<p class="nota">Exemplos reais de bairros que ficaram de fora:</p>' + render(ex_desc, transposed=False, colcls={"motivo": "alvo"}),
)
n_base_antes = len(base)
n_tirou = n_base_antes - len(modelagem)
bloco_b3 = caso(
    "Efeito do corte de bairros sobre as guias", "O merge final com <code>bairros</code> é <code>inner</code>: guias de bairro descartado saem.",
    f"−{fm(n_tirou)} guias ({pc(n_tirou / n_base_antes, 2)})",
    barra_tabela([("guias depois dos filtros e percentis", n_base_antes, n_base_antes), ("guias em bairros aprovados → base_modelagem", len(modelagem), n_base_antes)], cab=("etapa", "guias", "", "%")),
    "info",
)
secoes.append(passo(
    "s9", "9", "Agregação por bairro: <code>bairros.parquet</code>",
    f"Bairro é a unidade de análise: grande o bastante para ter crime e guias, pequeno o bastante para ter perfil próprio. {len(bairros)} bairros aprovados.",
    [bloco_b1, bloco_b2, bloco_b3],
))

# ───────────────────────── 10 · merge final ─────────────────────────
ORIGEM = {}
for c in ["sql", "logradouro", "numero", "bairro_itbi", "cep", "natureza_transacao", "valor_transacao", "data_transacao", "vvr", "proporcao_transmitida",
          "vvr_proporcional", "base_calculo", "valor_financiado", "situacao_sql", "area_terreno", "testada", "fracao_ideal", "area_construida", "uso_cod",
          "uso_desc", "padrao_cod", "padrao_desc", "acc"]:
    ORIGEM[c] = ("ITBI (planilha)", "mod" if c in ("sql", "logradouro", "bairro_itbi", "data_transacao") else "")
for c in ["setor_fiscal", "rua_key", "ano", "mes"]:
    ORIGEM[c] = ("ITBI (derivada na limpeza)", "novo")
for c in ["bairro", "lat", "lon"]:
    ORIGEM[c] = ("localiza() — passo 5", "novo")
ORIGEM["fator"] = ("IPCA — merge por mes (passo 7)", "novo")
for c in ["valor_real", "preco_m2_nominal", "preco_m2"]:
    ORIGEM[c] = ("deflação (passo 7)", "novo")
ORIGEM["idade"] = ("ano − acc válido (passo 6)", "novo")
ORIGEM["tipo"] = ("padrao_desc sem espaços", "novo")
for c in COLS_CRIME:
    ORIGEM[c] = ("perfil do bairro — merge por bairro (passo 9)", "novo")
assert set(ORIGEM) == set(base_pq.columns), set(base_pq.columns) ^ set(ORIGEM)

ra, rb2, rc = (modelagem[modelagem["sql"] == ALVOS[k]].iloc[0] for k in "ABC")
tr = pd.DataFrame([ra, rb2, rc])[list(base_pq.columns)].reset_index(drop=True)
linhas_dic = pd.DataFrame(
    [(c, ORIGEM[c][0], *(fmt(tr.at[i, c]) for i in range(3))) for c in base_pq.columns],
    columns=["coluna", "de onde vem", "guia A", "guia B", "guia C"],
)
cls_origem = {c: ORIGEM[c][1] for c in base_pq.columns}
# tabela com 3 guias reais lado a lado, colorida por origem
saida_final = ['<div class="rolagem"><table class="campos"><thead><tr><th>coluna</th><th>de onde vem</th><th>guia A</th><th>guia B</th><th>guia C</th></tr></thead><tbody>']
for _, r in linhas_dic.iterrows():
    k = cls_origem[r["coluna"]]
    saida_final.append(f'<tr><th>{e(r["coluna"])}</th><td class="orig">{e(r["de onde vem"])}</td>' + "".join(f'<td class="{k}">{r[g]}</td>' for g in ["guia A", "guia B", "guia C"]) + "</tr>")
saida_final.append("</tbody></table></div>")

merges = pd.DataFrame([
    ("localiza(): ITBI ⨝ setor_para_distrito", "setor_fiscal", "N → 1", fm(N), fm(N), "left"),
    ("localiza(): ITBI ⨝ voto por CEP", "cep_txt", "N → 1", fm(N), fm(N), "left"),
    ("localiza(): ITBI ⨝ dicionário", "(rua_key, distrito) → rua_key", "N → 1", fm(N), fm(N), "left"),
    ("base ⨝ IPCA", "mes", "N → 1", fm(n_antes_pct), fm(n_antes_pct), "left"),
    ("base ⨝ bairros", "bairro", "N → 1", fm(n_base_antes), fm(len(modelagem)), "inner (descarta bairro reprovado)"),
], columns=["merge", "chave", "cardinalidade", "linhas antes", "linhas depois", "tipo"])
bloco_m1 = caso(
    "Colunas de cada origem em 3 guias reais", "As três guias do rastreio (A, B, C) na linha final. A cor da célula mostra a origem: amarelo = coluna do ITBI que foi alterada na limpeza, verde = coluna criada pelo pipeline, branco = coluna da planilha passada adiante.",
    f"{fm(len(modelagem))} linhas × {base_pq.shape[1]} colunas", "".join(saida_final), "info",
)
bloco_m2 = caso(
    "Todos os merges e o que cada um fez com as linhas", "Contagens reais antes e depois. Só o último merge reduz linhas.", f"{len(merges)} merges",
    render(merges, transposed=False), "info",
)
secoes.append(passo(
    "s10", "10", "Base final: <code>base_modelagem.parquet</code>",
    "Cada linha é uma compra e venda residencial 2022–2026 com preço real, bairro, coordenada e o perfil criminal do bairro.",
    [bloco_m1, bloco_m2],
))

# ───────────────────────── 11 · limites ─────────────────────────
dup_itbi = int(itbi.duplicated().sum())
dup_base = int(modelagem.duplicated(subset=[c for c in modelagem.columns if c not in ("lat", "lon")]).sum())
dup_cr = int(crimes.drop(columns=["distrito", "grupo"]).duplicated().sum())
dup_ex = itbi[itbi.duplicated(keep=False)].sort_values(["sql", "data_transacao"]).head(2)[["sql", "logradouro", "numero", "data_transacao", "valor_transacao", "area_construida", "mes"]]
disc_cep_rua = loc[loc["bairro_cep"].notna() & loc["bairro_rua"].notna() & (loc["bairro_cep"] != loc["bairro_rua"])]
lat_falta = crimes.groupby("bairro", observed=True).agg(n=("lat", "size"), sem=("lat", lambda s: s.isna().mean()))
lat_falta = lat_falta[lat_falta["n"] >= 200]
lim = [
    caso("Distrito do ITBI é «o dominante do setor», não exato",
         f"Peso pelas guias, o distrito dominante cobre {pc(media_pond, 0)} da área do setor (passo 3). Usado só para desambiguar rua — mas uma rua num setor que cruza 4 distritos é desambiguada pelo distrito errado em parte dos casos.",
         f"{n_multi} de {len(setores)} setores", "", "limite"),
    caso("Linhas idênticas no ITBI não são removidas",
         "Nenhum passo faz <code>drop_duplicates</code>. As linhas abaixo têm todas as 27 colunas iguais; podem ser duas partes na mesma escritura ou duplicidade de exportação — não há como distinguir pelos dados.",
         f"{fm(dup_itbi)} linhas ({pc(dup_itbi / TOTAL_ITBI, 2)}) · {fm(dup_base)} na base final",
         render(dup_ex, transposed=False), "limite"),
    caso("Linhas idênticas nos crimes",
         "Sem número de BO na planilha lida, duas ocorrências diferentes no mesmo lugar/hora são indistinguíveis de uma duplicata. Mantidas.",
         f"{fm(dup_cr)} linhas ({pc(dup_cr / TOTAL_CRIMES, 2)})", "", "limite"),
    caso("CEP e rua discordam com frequência",
         "Quando as duas fontes têm resposta, divergem. O pipeline prioriza o CEP, mas o dicionário de ruas poderia estar certo em parte desses casos.",
         f"{fm(len(disc_cep_rua))} guias ({pc(len(disc_cep_rua) / N, 1)} do total)", "", "limite"),
    caso("Coordenada sem separador decimal é descartada, não corrigida",
         "Valores como <code>-2354897372</code> (= −23,54897372 × 10⁸) são recuperáveis dividindo por 10⁸, mas a janela de validação os descarta. Contagem na varredura bruta de todas as abas.",
         f"{fm(RAW['crime_cont'].get('lat_sem_virgula', 0))} crimes ({pc(RAW['crime_cont'].get('lat_sem_virgula', 0) / abas['sp'].sum(), 2)})", "", "limite"),
    caso("Coordenada ausente não é aleatória",
         f"A fração de crimes sem lat/lon por bairro (bairros com ≥ 200 crimes) vai de {pc(lat_falta['sem'].min())} a {pc(lat_falta['sem'].max())}. Por isso o pipeline usa a coordenada só como chave de rua/distrito, nunca para contar crimes por região.",
         f"{pc(crimes['lat'].isna().mean())} sem coordenada", "", "limite"),
]
secoes.append(passo(
    "s11", "11", "O que o pipeline não resolve",
    "Casos de borda reais que ficam como estão, cada um com o número do tamanho do efeito.",
    lim,
))

# ═════════════════════════ página ═════════════════════════
NAV = [("s0", "0 · Rastreio de 4 guias"), ("s1", "1 · Crimes"), ("s2", "2 · ITBI"), ("s3", "3 · Pontes geográficas"), ("s4", "4 · Dicionário de ruas"),
       ("s5", "5 · localiza()"), ("s6", "6 · Filtros"), ("s7", "7 · IPCA"), ("s8", "8 · Grupos de crime"), ("s9", "9 · Bairros"),
       ("s10", "10 · Base final"), ("s11", "11 · Limites")]
nav_html = "".join(f'<a href="#{i}">{t}</a>' for i, t in NAV)

CSS = """
:root{--bg:#f6f5f2;--card:#fff;--tinta:#16150f;--mudo:#6b6a62;--linha:#dddbd2;--azul:#2a78d6;--laranja:#eb6834;--aqua:#12936a;--code:#efede6;
--amarelo:#fff1b8;--amarelo-t:#5c4600;--verde:#d9f2e4;--verde-t:#0d5c3d;--verm:#fbdcd8;--verm-t:#8a1f12;--alvo:#ffe0cc;--alvo-b:#eb6834}
@media (prefers-color-scheme:dark){:root:not([data-theme=light]){--bg:#121211;--card:#1c1c1a;--tinta:#ecebe5;--mudo:#9a998f;--linha:#34332f;--azul:#6aa7f0;--laranja:#f08a5e;--aqua:#3fc79a;--code:#2a2a27;
--amarelo:#4a3d0a;--amarelo-t:#ffe27a;--verde:#143a2a;--verde-t:#7fe0b4;--verm:#4a1e19;--verm-t:#ff9a8c;--alvo:#47291a;--alvo-b:#f08a5e;color-scheme:dark}}
*{box-sizing:border-box}html{scroll-behavior:smooth;scroll-padding-top:12px}
body{margin:0;background:var(--bg);color:var(--tinta);font:14px/1.5 system-ui,-apple-system,"Segoe UI",sans-serif}
.layout{display:grid;grid-template-columns:200px minmax(0,1fr);gap:28px;max-width:1360px;margin:0 auto;padding:28px 16px 80px}
nav{position:sticky;top:16px;align-self:start;display:flex;flex-direction:column;gap:2px;font-size:13px}
nav b{font-size:11px;letter-spacing:.06em;text-transform:uppercase;color:var(--mudo);margin-bottom:6px}
nav a{color:var(--tinta);text-decoration:none;padding:4px 8px;border-radius:6px}nav a:hover{background:var(--code)}
main{min-width:0}
h1{font-size:26px;margin:0 0 6px}.sub-t{color:var(--mudo);max-width:820px;margin:0 0 18px}
code{background:var(--code);padding:1px 5px;border-radius:4px;font-size:12px;word-break:break-word}
.cod{background:var(--code);padding:8px 10px;border-radius:6px;font-size:12px;overflow-x:auto;margin:8px 0}
.legenda{display:flex;gap:10px;flex-wrap:wrap;margin:0 0 26px;font-size:12.5px}
.legenda span,.par.mod,.par.novo,.par.cai,.par.alvo{padding:1px 8px;border-radius:10px}
.passo{margin-bottom:46px}
.passo h2{font-size:19px;margin:0 0 4px;display:flex;align-items:center;gap:10px}
.passo h2 span{background:var(--tinta);color:var(--bg);min-width:26px;height:26px;padding:0 6px;border-radius:13px;display:grid;place-items:center;font-size:13px;flex:none}
.intro{margin:0 0 14px;color:var(--mudo);max-width:900px}
.caso{background:var(--card);border:1px solid var(--linha);border-left:4px solid var(--aqua);border-radius:8px;padding:12px 14px;margin:0 0 14px;min-width:0}
.caso.limite{border-left-color:var(--laranja)}.caso.info{border-left-color:var(--azul)}
.caso-cab{display:flex;align-items:baseline;gap:10px;flex-wrap:wrap}
.caso h4{margin:0;font-size:15px;flex:1 1 280px}
.selo{font-size:10.5px;letter-spacing:.05em;text-transform:uppercase;padding:1px 7px;border-radius:9px;background:var(--verde);color:var(--verde-t);font-weight:600}
.limite .selo{background:var(--alvo);color:var(--verm-t)}.info .selo{background:var(--code);color:var(--mudo)}
.n{font-size:12px;color:var(--mudo);font-variant-numeric:tabular-nums}
.regra{margin:6px 0 10px;color:var(--mudo);max-width:980px}
.nota{font-size:12.5px;color:var(--mudo);margin:10px 0 4px}
.rolagem{overflow-x:auto;margin:4px 0}
table{border-collapse:collapse;font-size:12px;width:100%}
th,td{padding:3px 8px;border-bottom:1px solid var(--linha);text-align:left;vertical-align:top}
thead th{color:var(--mudo);font-weight:600;white-space:nowrap}
.largo td,.campos td{font-family:ui-monospace,Consolas,monospace}
.largo td{white-space:nowrap}
.campos th{font-family:ui-monospace,Consolas,monospace;font-weight:600;color:var(--mudo);white-space:nowrap}
.campos td{max-width:300px;overflow-wrap:anywhere}
.campos .orig{font-family:system-ui,sans-serif;color:var(--mudo);white-space:normal}
.diff th.reg{text-align:center;border-left:2px solid var(--linha);color:var(--tinta)}
.diff td:nth-child(odd){border-left:2px solid var(--linha)}
td.mod,.par.mod{background:var(--amarelo);color:var(--amarelo-t)}
td.novo,.par.novo,tr.novo td{background:var(--verde);color:var(--verde-t)}
td.cai,.par.cai,tr.cai td{background:var(--verm);color:var(--verm-t);text-decoration:line-through}
td.alvo,.par.alvo,th.alvo{background:var(--alvo);box-shadow:inset 0 0 0 1px var(--alvo-b)}
td.vazio{color:var(--mudo);text-align:center}
.nulo{color:var(--mudo)}.txt{color:var(--azul)}
.num{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
.barra{width:34%;min-width:80px}.barra i{display:block;height:8px;background:var(--aqua);border-radius:2px}
.tempo th{width:170px;white-space:normal;font-weight:600;color:var(--tinta);font-size:12px}.tempo th small{color:var(--mudo);font-weight:400}
.tempo td{padding:5px 8px}
.par{display:inline-block;margin:2px 6px 2px 0;padding:1px 8px;border-radius:5px;background:var(--code);font-family:ui-monospace,Consolas,monospace;font-size:12px}
.par b{font-weight:600;color:var(--mudo);margin-right:4px}
.par.mod b,.par.novo b,.par.cai b{color:inherit;opacity:.75}
@media(max-width:900px){.layout{grid-template-columns:1fr}nav{display:none}}
"""

pagina = f"""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Pipeline SP — passo a passo</title><style>{CSS}</style></head><body>
<div class="layout"><nav><b>Passos</b>{nav_html}</nav><main>
<h1>Pipeline SP — crimes × ITBI, linha por linha</h1>
<p class="sub-t">Cada tabela é <b>lida dos arquivos do projeto</b> na hora de gerar esta página: nada digitado à mão, e o recálculo é conferido contra os parquets dos notebooks
<code>01_preparacao</code> e <code>02_eda</code>. Cada bloco é um caso de borda com exemplo real; <b>antes</b> é a célula como veio da planilha, <b>depois</b> é o que o pipeline guarda.</p>
<div class="legenda"><span class="par mod">amarelo = valor alterado</span><span class="par novo">verde = campo/valor criado</span><span class="par cai">vermelho = descartado</span><span class="par alvo">laranja = célula-problema</span>
<span class="par">∅ = vazio</span><span class="par">"texto" entre aspas = como veio da planilha</span></div>
{"".join(secoes)}
</main></div></body></html>"""
SAIDA.write_text(pagina, encoding="utf-8")
print(f"escrito {SAIDA} ({SAIDA.stat().st_size / 1024:,.0f} KB)")
